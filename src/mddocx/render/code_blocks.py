"""Code and diagram surface: fenced listings, highlighting, and Mermaid images."""

from __future__ import annotations

from typing import TYPE_CHECKING

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

from mddocx.ast.block import CodeBlock
from mddocx.diagnostics import Diagnostic, MddocxError
from mddocx.diagrams import MermaidRenderer, MermaidRenderError
from mddocx.ooxml.text import clean_xml_text

from . import citations, figures

if TYPE_CHECKING:
    from mddocx.render.renderer import DocxRenderer


def render_code_block(renderer: DocxRenderer, node: CodeBlock) -> None:
    if (node.language or "").strip().lower() == "mermaid" and renderer.config.mermaid.enabled:
        mermaid = MermaidRenderer(
            max_source_chars=renderer.config.mermaid.max_source_chars,
            max_nodes=renderer.config.mermaid.max_nodes,
            max_edges=renderer.config.mermaid.max_edges,
        )
        try:
            path = mermaid.render(node.code, renderer.resolver.temp_dir)
        except MermaidRenderError as exc:
            if renderer.config.mermaid.fallback == "error":
                raise MddocxError(Diagnostic("error", "DIAGRAM201", str(exc))) from exc
            renderer.reporter.warn(
                "DIAGRAM201",
                f"Mermaid kept as editable code: {exc}",
                getattr(node.source, "file", None),
                getattr(node.source, "line", None),
            )
            write_fenced_code(renderer, node)
        else:
            diagram_caption = (
                (node.caption or ("Diagram" if node.identifier else None))
                if renderer.config.references.captions
                else None
            )
            if diagram_caption and renderer.config.references.figure_caption_position == "above":
                citations.render_caption(renderer, "Figure", diagram_caption, node.identifier)
            p = renderer.document.add_paragraph(style=renderer._style("MD Normal"))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_together = True
            figures.add_image_path(
                renderer,
                p,
                path,
                diagram_caption or "Mermaid diagram",
                title=diagram_caption,
            )
            if diagram_caption and renderer.config.references.figure_caption_position == "below":
                citations.render_caption(renderer, "Figure", diagram_caption, node.identifier)
            renderer.reporter.info(
                "DIAGRAM101",
                "Rendered Mermaid diagram locally.",
                getattr(node.source, "file", None),
                getattr(node.source, "line", None),
            )
        return
    listing_caption = (
        (node.caption or ("Code listing" if node.identifier else None))
        if renderer.config.references.captions
        else None
    )
    if listing_caption and renderer.config.references.listing_caption_position == "above":
        citations.render_caption(renderer, "Listing", listing_caption, node.identifier)
    write_fenced_code(renderer, node)
    if listing_caption and renderer.config.references.listing_caption_position == "below":
        citations.render_caption(renderer, "Listing", listing_caption, node.identifier)


def write_fenced_code(renderer: DocxRenderer, node: CodeBlock) -> None:
    from .text_layout import render_text_layout

    if render_text_layout(renderer, node):
        return
    line_numbers = (
        renderer.config.code.line_numbers if node.line_numbers is None else node.line_numbers
    )
    show_label = (
        renderer.config.code.show_language_label
        if node.show_language_label is None
        else node.show_language_label
    )
    highlighted = set(renderer.config.code.highlight_lines) | set(node.highlight_lines)
    language = (node.language or "text").strip() or "text"
    if show_label and language.lower() not in {"text", "plain", "plaintext"}:
        lp = renderer.document.add_paragraph(style=renderer._style("MD Code Label"))
        lp.add_run(clean_xml_text(language))

    lexer = None
    pyg_style = None
    if renderer.config.code.syntax_highlighting:
        try:
            from pygments.lexers import get_lexer_by_name
            from pygments.styles import get_style_by_name

            lexer = get_lexer_by_name(language)
            pyg_style = get_style_by_name("friendly")
        except Exception:
            lexer = None
            pyg_style = None

    lines = node.code.splitlines() or [""]
    width = len(str(len(lines)))
    for line_no, raw_line in enumerate(lines, 1):
        p = renderer.document.add_paragraph(style=renderer._style("MD Code"))
        renderer._apply_planned_paragraph_layout(p, node)
        p.paragraph_format.space_before = 0
        p.paragraph_format.space_after = 0
        ppr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "FFF2CC" if line_no in highlighted else "F3F3F3")
        ppr.append(shd)
        if line_numbers:
            prefix = f"{line_no:>{width}}  "
            nr = p.add_run(prefix)
            renderer._configure_run(nr, prefix, code=True)
            nr.font.color.rgb = RGBColor(110, 110, 110)
        if lexer is None or pyg_style is None:
            text = clean_xml_text(raw_line)
            run = p.add_run(text)
            renderer._configure_run(run, text, code=True)
            continue
        try:
            from pygments import lex

            tokens = list(lex(raw_line, lexer))
        except Exception:
            tokens = []
        if not tokens:
            run = p.add_run(clean_xml_text(raw_line))
            renderer._configure_run(run, raw_line, code=True)
            continue
        for token_type, token_text in tokens:
            token_text = token_text.rstrip("\n")
            if not token_text:
                continue
            token_text = clean_xml_text(token_text)
            run = p.add_run(token_text)
            renderer._configure_run(run, token_text, code=True)
            info = pyg_style.style_for_token(token_type)
            color = info.get("color")
            if color and len(color) == 6:
                run.font.color.rgb = RGBColor.from_string(color.upper())
            run.bold = bool(info.get("bold"))
            run.italic = bool(info.get("italic"))
            if info.get("underline"):
                run.underline = True
