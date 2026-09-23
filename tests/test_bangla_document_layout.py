from io import BytesIO

from docx import Document
from docx.oxml.ns import qn

from mddocx import MarkdownWord, RenderConfig
from mddocx.parser import MarkdownParser
from mddocx.ast.block import CodeBlock, Paragraph
from mddocx.parser.compatibility import normalize_export_citations
from mddocx.ooxml.text import configure_run_fonts


def test_export_closing_fence_does_not_swallow_following_blocks():
    source = "Title[cite: 1]\n\n```text\nLiteral[cite: 2]\n```[cite: 1]\n\nAfter\n"
    nodes = MarkdownParser().parse(source).children
    assert len(nodes) == 3
    assert isinstance(nodes[1], CodeBlock)
    assert nodes[1].code == "Literal[cite: 2]"
    assert isinstance(nodes[2], Paragraph)
    assert normalize_export_citations("`[cite: 1]` [cite: 2]\n") == "`[cite: 1]` \n"


def test_keep_preserves_citations_but_repairs_fence():
    source = "```text\na\n```[cite: 1]\nAfter"
    result = normalize_export_citations(source, strip=False)
    assert "```\n[cite: 1]\nAfter" in result


def test_ruled_plain_text_table_is_editable_and_preserves_continuations():
    source = "```text\nFloor | East | West\n----------\n4 | Alice | Bob\n | 1125 | 1125\n----------\n3 | Common | Carol\n```"
    doc = Document(BytesIO(MarkdownWord().render_string(source)))
    assert len(doc.tables) == 1
    assert doc.tables[0].cell(1, 1).text == "Alice\n1125"
    assert doc.tables[0].cell(2, 2).text == "Carol"


def test_programming_fence_is_never_interpreted_as_layout():
    source = "```python\nFloor | East | West\n----------\n4 | Alice | Bob\n----------\n3 | Common | Carol\n```"
    doc = Document(BytesIO(MarkdownWord().render_string(source)))
    assert not doc.tables


def test_bangla_complex_font_does_not_enable_rtl():
    doc = Document()
    run = doc.add_paragraph().add_run("বাংলা")
    configure_run_fonts(run, run.text, "Arial", None, "Nirmala UI", False)
    assert run._r.rPr.rFonts.get(qn("w:cs")) == "Nirmala UI"
    assert run._r.rPr.find(qn("w:rtl")) is None


def test_bengali_document_theme():
    cfg = RenderConfig(theme="bengali-document")
    doc = Document(BytesIO(MarkdownWord(cfg).render_string("বাংলা")))
    assert doc.styles["MD Normal"].font.name == "Nirmala UI"


def test_form_has_native_photo_box_and_address_columns():
    source = "```text\nনাম : উদাহরণ ┌────┐\nপিতা : উদাহরণ │ ছবি │\nমাতা : উদাহরণ └────┘\nধর্ম : ইসলাম\nপেশা : শিক্ষক\nস্থায়ী | বর্তমান\nগ্রাম : ক | গ্রাম : খ\nজেলা : ক | জেলা : খ\n```"
    doc = Document(BytesIO(MarkdownWord().render_string(source)))
    assert len(doc.tables) == 2
    from bijoy2unicode.converter import Unicode

    def convert(value):
        return Unicode().convertUnicodeToBijoy(value + "  ")[:-2]

    assert doc.tables[0].cell(0, 1).tables[0].cell(0, 0).text == convert("ছবি")
    assert doc.tables[1].cell(1, 1).text == convert("গ্রাম") + " : " + convert("খ")
    assert doc.tables[0]._tbl.getnext().tag == qn("w:p")
    assert doc.tables[0]._tbl.getnext().getnext() is doc.tables[1]._tbl
    assert not any(p.style.name == "MD Code" for p in doc.paragraphs)


def test_layout_table_merges_preserve_stair_and_roof_geometry():
    import yaml

    spec = {
        "widths": [15, 37, 7, 41],
        "rows": [
            ["Floor", "East", "West", ""],
            ["5", "Roof", "", "Shakila"],
            ["4", "Taslima", "Stairs", "Farzana"],
            ["3", "Common", "", "Taslima"],
            ["2", "Common", "", "Mahi"],
            ["1", "Hasna", "", "Hasna"],
        ],
        "merges": [
            {"from": [0, 2], "to": [0, 3]},
            {"from": [1, 1], "to": [1, 2]},
            {"from": [2, 2], "to": [5, 2]},
        ],
    }
    source = "```layout-table\n" + yaml.safe_dump(spec) + "```"
    doc = Document(BytesIO(MarkdownWord().render_string(source)))
    table = doc.tables[0]
    assert len(table.rows) == 6
    assert len(table.columns) == 4
    assert table.cell(0, 2)._tc is table.cell(0, 3)._tc
    assert table.cell(1, 1)._tc is table.cell(1, 2)._tc
    assert table.cell(2, 2)._tc is table.cell(5, 2)._tc
    assert table.cell(5, 2).text == "Stairs"
    assert table.cell(4, 3).text == "Mahi"
    assert table.cell(2, 2)._tc.tcPr.vMerge.val == "restart"


def test_layout_table_rejects_data_loss_and_overlapping_merges():
    import pytest
    import yaml
    from mddocx.diagnostics import MddocxError

    for merges, rows in [
        ([{"from": [0, 0], "to": [1, 0]}], [["A", "B"], ["C", "D"]]),
        ([{"from": [0, 0], "to": [1, 0]}, {"from": [1, 0], "to": [1, 1]}], [["A", "B"], ["", ""]]),
        ([{"from": [0, 0], "to": [8, 0]}], [["A", "B"], ["", "D"]]),
    ]:
        source = "```layout-table\n" + yaml.safe_dump({"rows": rows, "merges": merges}) + "```"
        with pytest.raises(MddocxError, match="Invalid layout-table"):
            MarkdownWord().render_string(source)


def test_default_sutonny_encoding_preserves_english_and_word_structure():
    doc = Document(BytesIO(MarkdownWord().render_string("**বাংলা English 123**")))
    p = doc.paragraphs[0]
    assert p.text == "evsjv English 123"
    assert all(run.bold for run in p.runs)
    assert p.runs[0].font.name == "SutonnyMJ"
    assert p.runs[-1].font.name == "Aptos"


def test_sutonny_lafala_uses_supported_glyph_not_soft_hyphen():
    doc = Document(BytesIO(MarkdownWord().render_string("আল্লাহ")))
    assert doc.paragraphs[0].text == "Avjøvn"
    assert "\u00ad" not in doc.paragraphs[0].text


def test_isolated_bengali_prekar_after_other_script_avoids_bijoy_loop(monkeypatch):
    from bijoy2unicode.converter import Unicode

    original = Unicode.convertUnicodeToBijoy

    def guarded(self, value):
        assert value[0] not in "িেৈ"
        return original(self, value)

    monkeypatch.setattr(Unicode, "convertUnicodeToBijoy", guarded)
    source = "বাংলা " + " ".join(f"ல{mark}র" for mark in "িেৈ")
    doc = Document(BytesIO(MarkdownWord().render_string(source)))
    runs = doc.paragraphs[0].runs
    assert runs[0].text == "evsjv"
    assert [(r.text, r.font.name) for r in runs if r.font.name == "Nirmala UI"] == [
        (mark + "র", "Nirmala UI") for mark in "িেৈ"
    ]


def test_bangla_font_can_preserve_unicode_when_explicitly_disabled():
    cfg = RenderConfig()
    cfg.fonts.bengali = None
    doc = Document(BytesIO(MarkdownWord(cfg).render_string("বাংলা")))
    assert doc.paragraphs[0].text == "বাংলা"
