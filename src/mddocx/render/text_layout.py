"""Editable layouts for high-confidence plain-text document forms."""

from __future__ import annotations

import re
import yaml
from mddocx.diagnostics import Diagnostic, MddocxError
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from mddocx.ooxml.utils import set_table_fixed_layout, set_table_grid_widths


def render_text_layout(renderer, node) -> bool:
    structured = (node.language or "").lower() == "layout-table"
    if not structured and (node.language or "").lower() not in {
        "text",
        "plain",
        "plaintext",
        "form",
        "",
    }:
        return False
    lines = node.code.splitlines()
    # A box and paired address columns distinguish forms from code samples.
    form = (
        any("┌" in line and "┐" in line for line in lines)
        and sum(":" in line or "ঃ" in line for line in lines) >= 5
        and sum("|" in line for line in lines) >= 3
    )
    ruled = (
        sum(bool(re.fullmatch(r"\s*-{5,}\s*", line)) for line in lines) >= 2
        and sum(line.count("|") == 2 for line in lines) >= 3
    )
    if not structured and not form and not ruled:
        return False

    def table(rows, widths, borders=False):
        # Google Docs merges adjacent DOCX tables and inherits the first grid.
        # A small real paragraph keeps independent form/address grids intact.
        body = renderer.document._element.body
        previous = body[-2] if len(body) > 1 else None
        if previous is not None and previous.tag == qn("w:tbl"):
            spacer = renderer.document.add_paragraph()
            spacer.paragraph_format.space_before = Pt(0)
            spacer.paragraph_format.space_after = Pt(0)
            spacer.paragraph_format.line_spacing = Pt(1)
            spacer.add_run().font.size = Pt(1)
        result = renderer.document.add_table(rows=len(rows), cols=len(widths))
        result.autofit = False
        set_table_fixed_layout(result)
        set_table_grid_widths(result, [int(w * 56.6929133858) for w in widths])
        result.style = "Table Grid" if borders else "Normal Table"
        for col, width in zip(result.columns, widths):
            col.width = Mm(width)
        for row_index, (row, values) in enumerate(zip(result.rows, rows)):
            pr = row._tr.get_or_add_trPr()
            pr.append(OxmlElement("w:cantSplit"))
            for cell, value, width in zip(row.cells, values, widths):
                cell.width = Mm(width)
                p = cell.paragraphs[0]
                p.style = renderer._style("MD Normal")
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1
                run = p.add_run(value)
                renderer._configure_run(run, value)
                if borders:
                    p.paragraph_format.keep_with_next = row_index < len(rows) - 1
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return result

    section = renderer.document.sections[-1]
    available = (section.page_width - section.left_margin - section.right_margin) / 36000
    if structured:
        try:
            spec = yaml.safe_load(node.code)
            if not isinstance(spec, dict):
                raise ValueError("expected a mapping")
            rows = spec.get("rows")
            if (
                not isinstance(rows, list)
                or not rows
                or any(not isinstance(row, list) for row in rows)
            ):
                raise ValueError("rows must be a nonempty list of cell lists")
            cols = len(rows[0])
            if not cols or any(len(row) != cols for row in rows):
                raise ValueError("all rows must have the same number of cells")
            if any(not isinstance(cell, str) for row in rows for cell in row):
                raise ValueError("cells must be strings; use empty strings for merged placeholders")
            weights = spec.get("widths", [1] * cols)
            if (
                not isinstance(weights, list)
                or len(weights) != cols
                or any(
                    isinstance(w, bool) or not isinstance(w, (int, float)) or not 0 < w < 100000
                    for w in weights
                )
            ):
                raise ValueError("widths must contain one positive finite weight per column")
            merges = spec.get("merges", [])
            if not isinstance(merges, list):
                raise ValueError("merges must be a list")
            occupied = set()
            validated = []
            for merge in merges:
                if not isinstance(merge, dict):
                    raise ValueError("each merge must contain from and to coordinates")
                start, end = merge.get("from"), merge.get("to")
                if any(
                    not isinstance(pair, list)
                    or len(pair) != 2
                    or any(type(n) is not int for n in pair)
                    for pair in (start, end)
                ):
                    raise ValueError("merge coordinates must be [row, column] integer pairs")
                r1, c1 = start
                r2, c2 = end
                if not (0 <= r1 <= r2 < len(rows) and 0 <= c1 <= c2 < cols):
                    raise ValueError("merge rectangle is outside the table or reversed")
                rectangle = {(r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)}
                if occupied & rectangle:
                    raise ValueError("merge rectangles cannot overlap")
                if any(rows[r][c] for r, c in rectangle - {(r1, c1)}):
                    raise ValueError("covered merge cells must be empty to prevent data loss")
                occupied.update(rectangle)
                validated.append((r1, c1, r2, c2))
        except (ValueError, TypeError, yaml.YAMLError) as exc:
            raise MddocxError(
                Diagnostic(
                    "error",
                    "LAYOUT401",
                    f"Invalid layout-table: {exc}",
                    getattr(node.source, "file", None),
                    getattr(node.source, "line", None),
                )
            ) from exc
        result = table(rows, [available * w / sum(weights) for w in weights], True)
        for row in result.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r1, c1, r2, c2 in validated:
            value = rows[r1][c1]
            cell = result.cell(r1, c1).merge(result.cell(r2, c2))
            # Word merge concatenates placeholder paragraphs: retain one paragraph.
            for p in list(cell.paragraphs)[1:]:
                cell._tc.remove(p._p)
            p = cell.paragraphs[0]
            p.clear()
            renderer._configure_run(p.add_run(value), value)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return True
    if ruled:
        rows = []
        continuation = False
        for line in lines:
            if re.fullmatch(r"\s*-{5,}\s*", line):
                continuation = False
                continue
            if line.count("|") != 2:
                return False
            cells = [part.strip() for part in line.split("|")]
            if continuation and not cells[0]:
                rows[-1] = [
                    old + ("\n" if old and new else "") + new for old, new in zip(rows[-1], cells)
                ]
            else:
                rows.append(cells)
            continuation = True
        table(rows, [available * 0.16, available * 0.42, available * 0.42], True)
        return True

    fields, addresses, signatures = [], [], []
    for line in lines:
        line = re.split("[┌│└]", line, maxsplit=1)[0].strip()
        if not line:
            continue
        if "|" in line:
            addresses.append([part.strip() for part in line.split("|", 1)])
        elif re.match(r"\.{3,}", line):
            signatures.append(line)
        else:
            fields.append([line])
    outer = table([["", ""]], [available * 0.76, available * 0.24])
    left, photo = outer.rows[0].cells
    for index, values in enumerate(fields):
        p = left.paragraphs[0] if index == 0 else left.add_paragraph()
        p.style = renderer._style("MD Normal")
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1
        # Use tab stops rather than Unicode character counts for shaped Bangla.
        p.paragraph_format.tab_stops.add_tab_stop(Mm(34))
        value = re.sub(r"\s*:\s*", "\t: ", values[0], count=1)
        renderer._configure_run(p.add_run(value), value)
    p = photo.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    renderer._configure_run(p.add_run("ছবি"), "ছবি")
    # Draw an actual bordered photo box inside the right-hand layout cell.
    box = photo.add_table(rows=1, cols=1)
    box.cell(0, 0).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    bp = box.cell(0, 0).paragraphs[0]
    bp.style = renderer._style("MD Normal")
    renderer._configure_run(bp.add_run("ছবি"), "ছবি")
    box.autofit = False
    set_table_fixed_layout(box)
    set_table_grid_widths(box, [1361])
    box.cell(0, 0).width = Mm(24)
    box.rows[0].height = Mm(24)
    box.columns[0].width = Mm(24)
    box.style = "Table Grid"
    box.cell(0, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    photo._tc.remove(p._p)
    if addresses:
        table(addresses, [available / 2, available / 2])
    for value in signatures:
        p = renderer.document.add_paragraph(style=renderer._style("MD Normal"))
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(3)
        renderer._configure_run(p.add_run(value), value)
    return True
