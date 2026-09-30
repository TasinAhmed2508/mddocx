"""Multi-language output contract.

Markdown and the canonical AST stay Unicode. Script-specific font slots,
paragraph direction, and the optional Bengali legacy-font pass exist only in
the generated WordprocessingML package.
"""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document as WordDocument
from docx.oxml.ns import qn
from lxml import etree
import pytest

from mddocx import (
    AbstractConfig,
    FooterConfig,
    HeaderConfig,
    MarkdownWord,
    NotesConfig,
    RenderConfig,
    TitlePageConfig,
    inspect_docx_bytes,
)
from mddocx.cli import main
from mddocx.ooxml.text import (
    clean_xml_text,
    configure_run_fonts,
    has_cjk,
    has_complex_script,
    is_rtl_text,
)
from mddocx.parser import MarkdownParser

ROOT = Path(__file__).parents[1]
FIXTURE = ROOT / "tests" / "fixtures" / "multilingual" / "polyglot.md"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

COMPLEX_SCRIPT_SAMPLES = [
    "مرحبا",  # Arabic
    "עברית",  # Hebrew
    "हिन्दी",  # Devanagari
    "বাংলা",  # Bengali
    "ਪੰਜਾਬੀ",  # Gurmukhi
    "ગુજરાતી",  # Gujarati
    "ଓଡ଼ିଆ",  # Oriya
    "தமிழ்",  # Tamil
    "తెలుగు",  # Telugu
    "ಕನ್ನಡ",  # Kannada
    "മലയാളം",  # Malayalam
    "සිංහල",  # Sinhala
    "ไทย",  # Thai
    "ລາວ",  # Lao
    "བོད་ཡིག",  # Tibetan
    "မြန်မာ",  # Myanmar
    "አማርኛ",  # Ethiopic
    "ខ្មែរ",  # Khmer
]

LTR_SCRIPT_SAMPLES = [
    "English",
    "Русский",
    "Ελληνικά",
    "Türkçe",
    "Tiếng Việt",
    "中文",
    "日本語",
    "한국어",
    "12345",
    "🧪",
]

POLYGLOT_SAMPLES = [
    "English stays Latin",
    "Русский stays Cyrillic",
    "Ελληνικά stays Greek",
    "مرحبا بالعالم",
    "שלום עולם",
    "বাংলা ভাষা",
    "हिन्दी भाषा",
    "தமிழ் மொழி",
    "ไทย",
    "မြန်မာ",
    "ខ្មែរ",
    "中文",
    "日本語",
    "한국어",
]


def _q(tag: str) -> str:
    return f"{{{W}}}{tag}"


def _find(element, *tags: str):
    for tag in tags:
        if element is None:
            return None
        element = element.find(_q(tag))
    return element


def _part_xml(blob: bytes, part: str):
    with ZipFile(BytesIO(blob)) as package:
        assert package.testzip() is None
        return etree.fromstring(package.read(part))


def _document_xml(blob: bytes):
    return _part_xml(blob, "word/document.xml")


def _paragraphs(blob: bytes) -> list[dict]:
    result = []
    for paragraph in _document_xml(blob).iter(_q("p")):
        properties = paragraph.find(_q("pPr"))
        bidi = properties is not None and properties.find(_q("bidi")) is not None
        justify = None
        if properties is not None:
            alignment = properties.find(_q("jc"))
            justify = None if alignment is None else alignment.get(_q("val"))
        runs = []
        for run in paragraph.findall(_q("r")):
            rpr = run.find(_q("rPr"))
            fonts = None if rpr is None else rpr.find(_q("rFonts"))
            runs.append(
                {
                    "text": "".join(node.text or "" for node in run.iter(_q("t"))),
                    "ascii": None if fonts is None else fonts.get(_q("ascii")),
                    "east_asia": None if fonts is None else fonts.get(_q("eastAsia")),
                    "complex_script": None if fonts is None else fonts.get(_q("cs")),
                    "rtl": rpr is not None and rpr.find(_q("rtl")) is not None,
                }
            )
        result.append(
            {
                "text": "".join(node.text or "" for node in paragraph.iter(_q("t"))),
                "style": (
                    None
                    if properties is None or properties.find(_q("pStyle")) is None
                    else properties.find(_q("pStyle")).get(_q("val"))
                ),
                "bidi": bidi,
                "justify": justify,
                "runs": runs,
            }
        )
    return result


def _ast_text(node) -> str:
    parts = [
        value
        for attribute in ("text", "code", "source", "value")
        if isinstance((value := getattr(node, attribute, None)), str)
    ]
    for child in getattr(node, "children", None) or ():
        parts.append(_ast_text(child))
    for item in getattr(node, "items", None) or ():
        parts.append(_ast_text(item))
    return "".join(parts)


def _unicode_config() -> RenderConfig:
    config = RenderConfig(rtl="auto")
    config.fonts.bengali = None
    config.fonts.east_asia = "Microsoft YaHei"
    config.fonts.complex_script = "Nirmala UI"
    return config


@pytest.mark.parametrize("sample", COMPLEX_SCRIPT_SAMPLES)
def test_complex_script_detection_covers_word_writing_systems(sample: str):
    assert has_complex_script(sample)


@pytest.mark.parametrize("sample", LTR_SCRIPT_SAMPLES)
def test_plain_and_cjk_text_is_not_treated_as_complex_script(sample: str):
    assert not has_complex_script(sample)


@pytest.mark.parametrize(
    "sample",
    ["中文", "日本語", "한국어", "ｶﾀｶﾅ", "\U00020000", "中文 English"],
)
def test_cjk_detection_covers_east_asian_planes(sample: str):
    assert has_cjk(sample)


@pytest.mark.parametrize("sample", ["English", "Русский", "Ελληνικά", "مرحبا", "12345"])
def test_cjk_detection_rejects_other_scripts(sample: str):
    assert not has_cjk(sample)


@pytest.mark.parametrize(
    ("sample", "expected"),
    [
        ("مرحبا بالعالم", True),
        ("שלום עולם", True),
        ("فارسی ۱۲۳", True),
        ("Mixed English and مرحبا here", False),
        ("English only", False),
        ("12345", False),
        ("", False),
    ],
)
def test_rtl_detection_handles_bidi_text(sample: str, expected: bool):
    assert is_rtl_text(sample) is expected


def test_run_font_slots_are_script_specific():
    document = WordDocument()
    paragraph = document.add_paragraph()

    def add(text: str, rtl: bool = False):
        run = paragraph.add_run(text)
        configure_run_fonts(run, text, "Aptos", "Microsoft YaHei", "Nirmala UI", rtl)
        return run

    latin = add("English")
    chinese = add("中文")
    hindi = add("हिन्दी")
    arabic = add("مرحبا", rtl=True)

    assert latin._r.rPr.rFonts.get(qn("w:eastAsia")) is None
    assert latin._r.rPr.rFonts.get(qn("w:cs")) is None
    assert latin._r.rPr.find(qn("w:rtl")) is None

    assert chinese._r.rPr.rFonts.get(qn("w:eastAsia")) == "Microsoft YaHei"
    assert chinese._r.rPr.rFonts.get(qn("w:cs")) is None

    # Indic scripts are complex scripts but stay left-to-right.
    assert hindi._r.rPr.rFonts.get(qn("w:cs")) == "Nirmala UI"
    assert hindi._r.rPr.find(qn("w:rtl")) is None

    assert arabic._r.rPr.rFonts.get(qn("w:cs")) == "Nirmala UI"
    assert arabic._r.rPr.find(qn("w:rtl")) is not None


@pytest.mark.parametrize(
    ("mode", "text", "expected_rtl"),
    [
        ("auto", "مرحبا بالعالم", True),
        ("auto", "English only", False),
        ("auto", "Mixed English and مرحبا here", False),
        ("off", "مرحبا بالعالم", False),
        ("force", "English only", True),
    ],
)
def test_paragraph_direction_modes_control_bidi_and_alignment(
    mode: str, text: str, expected_rtl: bool
):
    config = RenderConfig(rtl=mode)
    config.fonts.bengali = None
    paragraph = _paragraphs(MarkdownWord(config).render_string(text))[0]

    assert paragraph["bidi"] is expected_rtl
    assert (paragraph["justify"] == "right") is expected_rtl


def test_cjk_runs_use_east_asia_font_only_when_needed():
    paragraphs = _paragraphs(MarkdownWord(_unicode_config()).render_string("English **中文** tail"))
    runs = {run["text"]: run for run in paragraphs[0]["runs"]}

    assert runs["中文"]["east_asia"] == "Microsoft YaHei"
    assert runs["English "]["east_asia"] is None
    assert all(not run["rtl"] for run in paragraphs[0]["runs"])


def test_indian_and_southeast_asian_scripts_use_complex_script_font():
    paragraph = _paragraphs(MarkdownWord(_unicode_config()).render_string("हिन्दी தமிழ் ไทย"))[0]
    (run,) = [run for run in paragraph["runs"] if run["text"].strip()]

    assert run["complex_script"] == "Nirmala UI"
    assert run["rtl"] is False


def test_bengali_boundary_preserves_unicode_ast_and_encodes_only_docx():
    source = "# বাংলা শিরোনাম\n\nবাংলা অনুচ্ছেদ English\n"
    parsed = MarkdownParser().parse(source)

    assert "বাংলা শিরোনাম" in _ast_text(parsed)

    unicode_document = WordDocument(BytesIO(MarkdownWord(_unicode_config()).render_string(source)))
    assert "বাংলা শিরোনাম" in [paragraph.text for paragraph in unicode_document.paragraphs]

    legacy_document = WordDocument(BytesIO(MarkdownWord().render_string(source)))
    texts = [paragraph.text for paragraph in legacy_document.paragraphs]

    assert "evsjv wk‡ivbvg" in texts
    assert "বাংলা" not in "".join(texts)
    assert "English" in "".join(texts)
    assert "বাংলা শিরোনাম" in _ast_text(MarkdownParser().parse(source))


def test_bangla_boundary_leaves_other_complex_scripts_as_unicode():
    source = "हिन्दी தமிழ் ไทย မြန်မာ ខ្មែរ العربية 中文"
    document = WordDocument(BytesIO(MarkdownWord().render_string(source)))
    text = "".join(paragraph.text for paragraph in document.paragraphs)

    for sample in ["हिन्दी", "தமிழ்", "ไทย", "မြန်မာ", "ខ្មែរ", "العربية", "中文"]:
        assert sample in text


def test_polyglot_fixture_keeps_every_script_and_native_structure():
    markdown = FIXTURE.read_text(encoding="utf-8")
    blob = MarkdownWord(_unicode_config()).render_string(markdown, base_dir=FIXTURE.parent)
    document = WordDocument(BytesIO(blob))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    inspection = inspect_docx_bytes(blob)

    for sample in POLYGLOT_SAMPLES:
        assert sample in text or sample in table_text

    assert "বাংলা" in table_text
    assert "三番目" in text
    assert any(paragraph.style.name == "MD Quote" for paragraph in document.paragraphs)
    assert "# " not in text and "**" not in text and "|" not in text
    assert inspection.headings >= 5
    assert inspection.tables == 1
    assert inspection.native_list_paragraphs >= 3
    assert inspection.equations >= 1
    assert inspection.ok


def test_cli_script_font_flags_reach_the_document(tmp_path: Path, capsys):
    output = tmp_path / "polyglot.docx"

    assert (
        main(
            [
                str(FIXTURE),
                "-o",
                str(output),
                "--east-asia-font",
                "Microsoft YaHei",
                "--complex-script-font",
                "Nirmala UI",
            ]
        )
        == 0
    )
    capsys.readouterr()
    root = _document_xml(output.read_bytes())
    fonts = list(root.iter(_q("rFonts")))

    assert any(element.get(_q("eastAsia")) == "Microsoft YaHei" for element in fonts)
    assert any(element.get(_q("cs")) == "Nirmala UI" for element in fonts)


def test_rtl_reaches_title_page_abstract_headers_and_footers():
    config = RenderConfig(
        rtl="auto",
        title="تقرير سنوي",
        author="مؤلف عربي",
        title_page=TitlePageConfig(enabled=True, subtitle="تقرير فرعي"),
        abstract=AbstractConfig(text="الملخص العربي"),
        header=HeaderConfig(text="الفصل الأول"),
        footer=FooterConfig(text="שלום עולם"),
    )
    config.fonts.bengali = None
    config.fonts.complex_script = "Nirmala UI"
    blob = MarkdownWord(config).render_string("# Contents\n\nBody text.\n")

    header = _part_xml(blob, "word/header1.xml")
    footer = _part_xml(blob, "word/footer1.xml")
    document = _document_xml(blob)

    for part in (header, footer):
        paragraphs = list(part.iter(_q("p")))
        assert any(_find(paragraph, "pPr") is not None for paragraph in paragraphs)
        assert any(_find(paragraph, "pPr", "bidi") is not None for paragraph in paragraphs)
        assert all(
            _find(run, "rPr", "rtl") is not None
            for paragraph in paragraphs
            for run in paragraph.findall(_q("r"))
            if "".join(node.text or "" for node in run.iter(_q("t"))).strip()
        )

    # Centered title-page lines keep their explicit alignment; body-style text
    # is right-aligned by the direction pass.
    expectations = {
        "تقرير سنوي": "center",
        "تقرير فرعي": "center",
        "مؤلف عربي": "center",
        "الملخص العربي": "right",
    }
    for sample, expected_justify in expectations.items():
        paragraph = next(
            item
            for item in document.iter(_q("p"))
            if sample in "".join(node.text or "" for node in item.iter(_q("t")))
        )
        assert _find(paragraph, "pPr", "bidi") is not None
        assert _find(paragraph, "pPr", "jc").get(_q("val")) == expected_justify
        for run in paragraph.findall(_q("r")):
            fonts = _find(run, "rPr", "rFonts")
            assert fonts is not None and fonts.get(_q("cs")) == "Nirmala UI"
            assert _find(run, "rPr", "rtl") is not None


def test_rtl_note_text_gets_direction_in_footnotes_and_endnotes():
    source = "نص عربي[^1]\n\n[^1]: ملاحظة عربية\n"
    for style, part, note_text in (
        ("footnote", "word/footnotes.xml", "ملاحظة عربية"),
        ("endnote", "word/endnotes.xml", "ملاحظة عربية"),
    ):
        config = RenderConfig(rtl="auto", notes=NotesConfig(style=style))
        config.fonts.bengali = None
        config.fonts.complex_script = "Nirmala UI"
        blob = MarkdownWord(config).render_string(source)
        root = _part_xml(blob, part)
        paragraph = next(
            item
            for item in root.iter(_q("p"))
            if note_text in "".join(node.text or "" for node in item.iter(_q("t")))
        )

        assert _find(paragraph, "pPr", "bidi") is not None
        assert _find(paragraph, "pPr", "jc").get(_q("val")) == "right"
        run = next(
            item
            for item in paragraph.findall(_q("r"))
            if note_text in "".join(node.text or "" for node in item.iter(_q("t")))
        )
        assert _find(run, "rPr", "rFonts").get(_q("cs")) == "Nirmala UI"
        assert _find(run, "rPr", "rtl") is not None


def test_bangla_boundary_encodes_headers_tables_and_notes():
    config = RenderConfig(
        header=HeaderConfig(text="বাংলা শিরোনাম"),
        footer=FooterConfig(text="বাংলা পাদ"),
    )
    markdown = "| ভাষা |\n| --- |\n| বাংলা |\n\nটীকা[^1]\n\n[^1]: বাংলা পাদটীকা\n"
    blob = MarkdownWord(config).render_string(markdown)

    for part in ("word/header1.xml", "word/footer1.xml", "word/document.xml", "word/footnotes.xml"):
        text = "".join(node.text or "" for node in _part_xml(blob, part).iter(_q("t")))
        assert "evsjv" in text
        assert not any("\u0980" <= character <= "\u09ff" for character in text)


def test_clean_xml_text_keeps_astral_scripts_and_drops_invalid_controls():
    cleaned = clean_xml_text("🧪中文\U0002a6d6\x00\x1f\t\n")

    assert cleaned == "🧪中文\U0002a6d6\t\n"
