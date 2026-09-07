from __future__ import annotations

import re

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def add_hyperlink(
    paragraph,
    text: str,
    url: str,
    bold: bool = False,
    italic: bool = False,
    strike: bool = False,
    font_name: str | None = None,
    rtl: bool = False,
):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1"); r_pr.append(color)
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); r_pr.append(underline)
    if font_name:
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), font_name); fonts.set(qn("w:hAnsi"), font_name)
        if rtl: fonts.set(qn("w:cs"), font_name)
        r_pr.append(fonts)
    if bold: r_pr.append(OxmlElement("w:b"))
    if italic: r_pr.append(OxmlElement("w:i"))
    if strike: r_pr.append(OxmlElement("w:strike"))
    if rtl:
        rtl_el = OxmlElement("w:rtl"); rtl_el.set(qn("w:val"), "1"); r_pr.append(rtl_el)
    run.append(r_pr)
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text; run.append(t)
    hyperlink.append(run); paragraph._p.append(hyperlink)
    return hyperlink


def add_field(paragraph, instruction: str, placeholder: str | None = None, dirty: bool = True) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    if dirty:
        begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")

    r1 = OxmlElement("w:r"); r1.append(begin)
    r2 = OxmlElement("w:r"); r2.append(instr)
    r3 = OxmlElement("w:r"); r3.append(separate)
    paragraph._p.extend([r1, r2, r3])
    if placeholder:
        paragraph.add_run(placeholder)
    r4 = OxmlElement("w:r"); r4.append(end); paragraph._p.append(r4)


def add_page_number(paragraph) -> None:
    add_field(paragraph, " PAGE ", "1")


def add_toc(paragraph, min_level: int = 1, max_level: int = 3) -> None:
    minimum = max(1, min(9, int(min_level)))
    maximum = max(minimum, min(9, int(max_level)))
    add_field(paragraph, f' TOC \\o "{minimum}-{maximum}" \\h \\z \\u ', "Update this field in Word")


def add_style_ref(paragraph, style_name: str = "MD Heading 1") -> None:
    add_field(paragraph, f' STYLEREF "{style_name}" ', "Section")


def bookmark_name(text: str, fallback_index: int) -> str:
    safe = re.sub(r"[^A-Za-z0-9_]", "_", text).strip("_")
    if not safe or not safe[0].isalpha():
        safe = f"Heading_{fallback_index}_{safe}"
    return safe[:40]


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_hyperlink(paragraph, text: str, anchor: str):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1"); r_pr.append(color)
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); r_pr.append(underline)
    run.append(r_pr)
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text; run.append(t)
    hyperlink.append(run); paragraph._p.append(hyperlink)
    return hyperlink


def add_ref_field(paragraph, bookmark: str, placeholder: str, hyperlink: bool = True) -> None:
    switch = " \\h" if hyperlink else ""
    add_field(paragraph, f" REF {bookmark}{switch} ", placeholder)


def add_seq_field(paragraph, label: str, placeholder: str, bookmark: str | None = None, bookmark_id: int | None = None) -> None:
    if bookmark and bookmark_id is not None:
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(bookmark_id)); start.set(qn("w:name"), bookmark)
        paragraph._p.append(start)
    add_field(paragraph, f" SEQ {label} \\* ARABIC ", placeholder)
    if bookmark and bookmark_id is not None:
        end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(bookmark_id)); paragraph._p.append(end)


def add_hidden_field(paragraph, instruction: str, placeholder: str = "") -> None:
    """Insert a Word field whose result is hidden but whose sequence side effect remains."""
    def hidden_run(child):
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        vanish = OxmlElement("w:vanish"); vanish.set(qn("w:val"), "1"); rpr.append(vanish)
        run.append(rpr); run.append(child); return run
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin"); begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = instruction
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    paragraph._p.extend([hidden_run(begin), hidden_run(instr), hidden_run(separate)])
    if placeholder:
        t = OxmlElement("w:t"); t.text = placeholder
        paragraph._p.append(hidden_run(t))
    paragraph._p.append(hidden_run(end))


def add_section_equation_number(paragraph, section_placeholder: str, equation_placeholder: str, bookmark: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id)); start.set(qn("w:name"), bookmark)
    paragraph._p.append(start)
    add_field(paragraph, " SEQ Section \\c ", section_placeholder)
    paragraph.add_run(".")
    add_field(paragraph, " SEQ Equation \\* ARABIC ", equation_placeholder)
    end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(bookmark_id)); paragraph._p.append(end)


def add_num_pages(paragraph) -> None:
    add_field(paragraph, " NUMPAGES ", "1")


def add_page_x_of_y(paragraph) -> None:
    paragraph.add_run("Page ")
    add_page_number(paragraph)
    paragraph.add_run(" of ")
    add_num_pages(paragraph)


def add_named_field(paragraph, name: str, placeholder: str = "") -> None:
    name = name.upper().strip()
    instructions = {
        "DATE": r' DATE \@ "MMMM d, yyyy" ',
        "CREATEDATE": r' CREATEDATE \@ "MMMM d, yyyy" ',
        "SAVEDATE": r' SAVEDATE \@ "MMMM d, yyyy" ',
        "AUTHOR": " AUTHOR ",
        "NUMPAGES": " NUMPAGES ",
        "SECTION": " SECTION ",
        "SECTIONPAGES": " SECTIONPAGES ",
        "FILENAME": r" FILENAME \p ",
        "TITLE": ' DOCPROPERTY "Title" ',
        "SUBJECT": ' DOCPROPERTY "Subject" ',
    }
    instruction = instructions.get(name)
    if instruction is None:
        if name.startswith("DOCPROPERTY:"):
            prop = name.split(":", 1)[1].strip().replace('"', "")
            instruction = f' DOCPROPERTY "{prop}" '
        else:
            return
    add_field(paragraph, instruction, placeholder)


_FIELD_TOKEN_RE = re.compile(r"\{(PAGE|NUMPAGES|DATE|CREATEDATE|SAVEDATE|AUTHOR|SECTION|SECTIONPAGES|FILENAME|TITLE|SUBJECT|DOCPROPERTY:[A-Za-z0-9 _.-]+)\}", re.I)


def render_field_template(paragraph, text: str, placeholders: dict[str, str] | None = None) -> None:
    """Render a small safe field-template language into native Word fields."""
    placeholders = {k.upper(): v for k, v in (placeholders or {}).items()}
    pos = 0
    for match in _FIELD_TOKEN_RE.finditer(text or ""):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        name = match.group(1).upper()
        if name == "PAGE":
            add_page_number(paragraph)
        else:
            add_named_field(paragraph, name, placeholders.get(name, ""))
        pos = match.end()
    if pos < len(text or ""):
        paragraph.add_run(text[pos:])


def add_chapter_seq_field(
    paragraph,
    label: str,
    chapter_placeholder: str,
    item_placeholder: str,
    bookmark: str | None = None,
    bookmark_id: int | None = None,
) -> None:
    """Emit Word-standard chapter-scoped caption numbering using STYLEREF + SEQ."""
    if bookmark and bookmark_id is not None:
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(bookmark_id)); start.set(qn("w:name"), bookmark)
        paragraph._p.append(start)
    add_field(paragraph, ' STYLEREF "MD Heading 1" \\n ', chapter_placeholder)
    paragraph.add_run(".")
    switch = r"\r 1 \s 1" if str(item_placeholder) == "1" else r"\s 1"
    add_field(paragraph, f" SEQ {label} \\* ARABIC {switch} ", item_placeholder)
    if bookmark and bookmark_id is not None:
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(bookmark_id)); paragraph._p.append(end)


def add_section_caption_number(
    paragraph,
    label: str,
    section_placeholder: str,
    item_placeholder: str,
    bookmark: str,
    bookmark_id: int,
) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id)); start.set(qn("w:name"), bookmark)
    paragraph._p.append(start)
    add_field(paragraph, " SEQ Section \\c ", section_placeholder)
    paragraph.add_run(".")
    add_field(paragraph, f" SEQ {label} \\* ARABIC ", item_placeholder)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id)); paragraph._p.append(end)
