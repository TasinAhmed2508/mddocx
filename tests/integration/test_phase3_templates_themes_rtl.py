from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt

from mddocx import FontConfig, MarkdownWord, RenderConfig, render_string
from mddocx.ast.base import Document as AstDocument, Node


def _zip_text(blob: bytes, name: str) -> str:
    with ZipFile(BytesIO(blob)) as z:
        return z.read(name).decode("utf-8")


def test_docx_template_preserves_body_styles_and_page_setup(tmp_path: Path):
    template = tmp_path / "template.docx"
    doc = Document()
    doc.sections[0].orientation = WD_ORIENT.LANDSCAPE
    doc.sections[0].page_width, doc.sections[0].page_height = doc.sections[0].page_height, doc.sections[0].page_width
    style = doc.styles.add_style("MD Normal", WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Courier New"
    style.font.size = Pt(13)
    doc.add_paragraph("Template cover")
    doc.save(template)

    blob = render_string("# Added\n\nBody", template=template)
    reopened = Document(BytesIO(blob))
    text = "\n".join(p.text for p in reopened.paragraphs)
    assert "Template cover" in text and "Added" in text and "Body" in text
    assert reopened.sections[0].page_width > reopened.sections[0].page_height
    assert reopened.styles["MD Normal"].font.name == "Courier New"
    assert round(reopened.styles["MD Normal"].font.size.pt) == 13


def test_theme_custom_fonts_unicode_and_rtl_emit_native_word_properties():
    cfg = RenderConfig(
        theme="academic",
        fonts=FontConfig(body="Liberation Serif", headings="Liberation Serif", east_asia="Noto Sans CJK SC", complex_script="Noto Naskh Arabic"),
        rtl="auto",
    )
    blob = render_string("# عنوان\n\nمرحبا بالعالم — 中文内容 — café 🚀", config=cfg)
    Document(BytesIO(blob))
    styles = _zip_text(blob, "word/styles.xml")
    xml = _zip_text(blob, "word/document.xml")
    assert "Liberation Serif" in styles
    assert "Noto Sans CJK SC" in styles
    assert "Noto Naskh Arabic" in styles
    assert "<w:bidi" in xml and "<w:rtl" in xml
    assert "中文内容" in xml and "café" in xml and "🚀" in xml


def test_invalid_xml_control_characters_are_sanitized():
    blob = render_string("Text before \x01 text after")
    xml = _zip_text(blob, "word/document.xml")
    assert "\x01" not in xml
    assert "Text before " in xml and " text after" in xml


@dataclass(slots=True)
class Callout(Node):
    text: str = ""


class CalloutExtension:
    def parse_block(self, parser, tokens, index):
        if tokens[index].type == "hr":
            return Callout(text="parsed extension token"), index + 1
        return None

    def render_block(self, renderer, node):
        if not isinstance(node, Callout):
            return False
        p = renderer.document.add_paragraph(style="MD Quote")
        p.add_run("CALL OUT: " + node.text)
        return True


def test_custom_ast_node_can_be_rendered_by_extension():
    cfg = RenderConfig(extensions=(CalloutExtension(),))
    converter = MarkdownWord(cfg)
    blob = converter.render_ast(AstDocument(children=[Callout(text="custom node")]))
    doc = Document(BytesIO(blob))
    assert any(p.text == "CALL OUT: custom node" for p in doc.paragraphs)


def test_extension_can_convert_parser_tokens_into_custom_ast():
    cfg = RenderConfig(extensions=(CalloutExtension(),))
    blob = render_string("---", config=cfg)
    doc = Document(BytesIO(blob))
    assert any(p.text == "CALL OUT: parsed extension token" for p in doc.paragraphs)


def test_auto_landscape_keeps_template_custom_page_dimensions(tmp_path: Path):
    from docx.shared import Mm
    from mddocx import TableConfig

    template = tmp_path / "custom-size.docx"
    doc = Document()
    doc.sections[0].page_width = Mm(180)
    doc.sections[0].page_height = Mm(250)
    doc.save(template)

    header = "| " + " | ".join(f"C{i}" for i in range(7)) + " |"
    sep = "|" + "|".join("---" for _ in range(7)) + "|"
    row = "| " + " | ".join("x" * 20 for _ in range(7)) + " |"
    cfg = RenderConfig(template=template, table=TableConfig(auto_landscape=True))
    blob = render_string(header + "\n" + sep + "\n" + row, config=cfg)
    reopened = Document(BytesIO(blob))
    assert len(reopened.sections) == 3
    dims = [(round(s.page_width.mm), round(s.page_height.mm)) for s in reopened.sections]
    assert dims == [(180, 250), (250, 180), (180, 250)]
