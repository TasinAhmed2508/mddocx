from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document as WordDocument
from lxml import etree
from PIL import Image
import pytest

from mddocx import (
    MarkdownWord,
    RenderConfig,
    TableConfig,
    audit_docx_accessibility_bytes,
    inspect_docx_bytes,
)
from mddocx.cli import main
from mddocx.diagnostics import MddocxError


ROOT = Path(__file__).parents[1]
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
}


def _document_xml(blob: bytes):
    with ZipFile(BytesIO(blob)) as package:
        assert package.testzip() is None
        return etree.fromstring(package.read("word/document.xml"))


def test_machine_readable_feature_matrix_has_live_evidence():
    matrix = json.loads((ROOT / "docs" / "feature-matrix.json").read_text(encoding="utf-8"))

    assert matrix["schema"] == "mddocx-feature-matrix"
    assert matrix["version"] == 1
    assert matrix["fidelity_dimensions"] == [
        "semantic",
        "native_editability",
        "presentation",
    ]
    ids = [feature["id"] for feature in matrix["features"]]
    assert len(ids) == len(set(ids)) >= 12
    required = {"syntax", "ast", "ooxml", "malformed_diagnostic", "word365", "libreoffice"}
    for feature in matrix["features"]:
        assert required.issubset(feature)
        assert feature["tests"]
        assert all((ROOT / test).exists() for test in feature["tests"])


def test_docx_round_trip_reopens_and_retains_semantic_structure():
    markdown = "# Heading\n\nParagraph with **bold**.\n\n1. First\n2. Second"
    blob = MarkdownWord().render_string(markdown)
    reopened = WordDocument(BytesIO(blob))
    inspection = inspect_docx_bytes(blob)

    assert reopened.paragraphs[0].text == "Heading"
    assert reopened.paragraphs[1].text == "Paragraph with bold."
    assert inspection.headings == 1
    assert inspection.native_list_paragraphs == 2
    assert inspection.ok


def test_native_image_has_alternative_text_and_accessibility_passes(tmp_path: Path):
    image_path = tmp_path / "sample.png"
    Image.new("RGB", (24, 16), color=(30, 80, 140)).save(image_path)
    blob = MarkdownWord().render_string(
        "# Illustrated\n\n![Blue sample](sample.png)", base_dir=tmp_path
    )
    root = _document_xml(blob)

    assert root.xpath(".//w:drawing", namespaces=NS)
    assert root.xpath(".//wp:docPr[@descr='Blue sample']", namespaces=NS)
    assert audit_docx_accessibility_bytes(blob).images_missing_alt == 0


def test_unicode_xml_safety_and_rtl_properties_are_preserved():
    blob = MarkdownWord(RenderConfig(rtl="auto")).render_string(
        "# Unicode\n\nمرحبا بالعالم — বাংলা — emoji 🧪\x00"
    )
    root = _document_xml(blob)
    text = root.xpath("string(.//w:body)", namespaces=NS)

    assert "مرحبا بالعالم" in text
    assert "বাংলা" in text
    assert "🧪" in text
    assert "\x00" not in text
    assert root.xpath(".//w:pPr/w:bidi", namespaces=NS)


def test_invalid_template_paths_have_stable_actionable_diagnostics(tmp_path: Path):
    missing = tmp_path / "missing.docx"
    with pytest.raises(MddocxError, match="template not found") as missing_error:
        MarkdownWord(RenderConfig(template=missing)).render_string("# Report")
    assert missing_error.value.diagnostic.code == "CONFIG301"

    wrong = tmp_path / "template.txt"
    wrong.write_text("not a package", encoding="utf-8")
    with pytest.raises(MddocxError, match="must be a .docx") as wrong_error:
        MarkdownWord(RenderConfig(template=wrong)).render_string("# Report")
    assert wrong_error.value.diagnostic.code == "CONFIG302"


def test_cli_exit_codes_and_paths_with_spaces(tmp_path: Path, capsys):
    source = tmp_path / "report with spaces.md"
    source.write_text("# Valid\n\nInline \\(x^2\\).", encoding="utf-8")

    assert main([str(source), "--check"]) == 0
    assert str(source) in capsys.readouterr().out

    broken = tmp_path / "unsupported math.md"
    broken.write_text("$$\\inventedmacro{x}$$", encoding="utf-8")
    assert main(["math-check", str(broken), "--strict"]) == 3
    assert "Fallbacks: 1" in capsys.readouterr().out


def test_plugin_transform_failures_are_isolated_with_stable_diagnostics():
    class RaisingExtension:
        def transform_document(self, _document):
            raise RuntimeError("private plugin detail")

    class InvalidExtension:
        def transform_document(self, _document):
            return "not an AST"

    with pytest.raises(MddocxError) as raised:
        MarkdownWord(RenderConfig(extensions=(RaisingExtension(),))).render_string("# Input")
    assert raised.value.diagnostic.code == "PLUGIN405"
    assert "RaisingExtension" in raised.value.diagnostic.message

    with pytest.raises(MddocxError) as invalid:
        MarkdownWord(RenderConfig(extensions=(InvalidExtension(),))).render_string("# Input")
    assert invalid.value.diagnostic.code == "PLUGIN406"
    assert "InvalidExtension" in invalid.value.diagnostic.message


def test_invalid_configuration_fails_at_the_compiler_boundary():
    config = RenderConfig(table=TableConfig(min_column_width_mm=50, max_column_width_mm=10))

    with pytest.raises(MddocxError) as error:
        MarkdownWord(config)

    assert error.value.diagnostic.code == "CONFIG401"
    assert "Maximum table column width" in error.value.diagnostic.message
