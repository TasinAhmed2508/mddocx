"""Reference surfaces: caption labels/fields, cross-references, and bibliography.

All three read the shared ``ReferenceRegistry`` and emit Word fields, so they own
the numbered-reference presentation rather than the surrounding block structure.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

from docx.shared import Mm

from mddocx.ast.inline import CrossReference
from mddocx.ooxml.fields import (
    add_bookmark,
    add_chapter_seq_field,
    add_hyperlink,
    add_ref_field,
    add_seq_field,
)
from mddocx.ooxml.text import clean_xml_text
from mddocx.references import ReferenceRegistry

if TYPE_CHECKING:
    from mddocx.render.renderer import DocxRenderer


def display_label(renderer: DocxRenderer, kind: str) -> str:
    return {
        "Figure": renderer.config.references.figure_label,
        "Table": renderer.config.references.table_label,
        "Equation": renderer.config.references.equation_label,
        "Listing": renderer.config.references.listing_label,
    }.get(kind, kind)


def render_caption(
    renderer: DocxRenderer, label: str, caption: str, identifier: str | None
) -> None:
    p = renderer.document.add_paragraph(style=renderer._style("MD Caption"))
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = True
    target = renderer.references.get(identifier) if identifier and renderer.references else None
    p.add_run(display_label(renderer, label) + " ")
    if target and target.number is not None:
        if (
            renderer.config.references.caption_number_format == "section"
            and isinstance(target.number, str)
            and "." in target.number
        ):
            section_no, item_no = target.number.split(".", 1)
            add_chapter_seq_field(
                p, label, section_no, item_no, target.bookmark, renderer._bookmark_id
            )
        else:
            add_seq_field(p, label, str(target.number), target.bookmark, renderer._bookmark_id)
        renderer._bookmark_id += 1
    else:
        add_seq_field(p, label, "1")
    if caption:
        p.add_run(" — " + clean_xml_text(caption))
        renderer._apply_text_policy(p, caption)
    else:
        renderer._apply_text_policy(p, label)


def render_cross_reference(renderer: DocxRenderer, paragraph, node: CrossReference) -> None:
    target = renderer.references.get(node.target) if renderer.references else None
    if not target:
        text = (node.prefix + " " if node.prefix else "") + "@" + node.target
        r = paragraph.add_run(text)
        renderer._configure_run(r, text)
        renderer.reporter.warn(
            "REF202",
            f"Unresolved cross-reference: {node.target}",
            getattr(node.source, "file", None),
            getattr(node.source, "line", None),
        )
        return
    prefix = node.prefix
    if not renderer.config.references.enabled:
        if (
            prefix is None
            and renderer.config.references.include_prefix_in_crossrefs
            and target.kind != "Section"
        ):
            prefix = display_label(renderer, target.kind)
        if target.kind == "Section":
            value = target.title or node.target
            if prefix and value.lower().startswith(prefix.lower()):
                prefix = None
        else:
            value = str(target.number or "?")
        text = ((prefix + " ") if prefix else "") + value
        r = paragraph.add_run(text)
        renderer._configure_run(r, text)
        return
    if (
        prefix is None
        and renderer.config.references.include_prefix_in_crossrefs
        and target.kind != "Section"
    ):
        prefix = display_label(renderer, target.kind)
    if (
        target.kind == "Section"
        and prefix
        and (target.title or "").lower().startswith(prefix.lower())
    ):
        prefix = None
    if prefix:
        r = paragraph.add_run(prefix + " ")
        renderer._configure_run(r, r.text)
    if target.kind == "Section":
        add_ref_field(paragraph, target.bookmark, target.title or node.target)
    else:
        add_ref_field(paragraph, target.bookmark, str(target.number or "?"))


def render_bibliography(renderer: DocxRenderer) -> None:
    if renderer._bibliography_rendered:
        return
    renderer._bibliography_rendered = True
    title = clean_xml_text(renderer.config.citations.bibliography_title)
    last = renderer.document.paragraphs[-1] if renderer.document.paragraphs else None
    has_explicit_heading = bool(
        last
        and last.text.strip().casefold() == title.strip().casefold()
        and last.style is not None
        and (last.style.name or "").startswith(("MD Heading", "Heading"))
    )
    if not has_explicit_heading:
        p = renderer.document.add_paragraph(style=renderer._style("MD Heading 1"))
        p.add_run(title)
        renderer._apply_text_policy(p, title)
    cited = set(renderer._cited_keys)
    include_all = getattr(renderer.config.citations, "bibliography_include", "cited") == "all"
    keys = None if include_all or not cited else renderer._cited_keys
    entries = renderer.bibliography.formatted_segments(renderer._citation_style, keys)
    for key, segments in entries:
        bp = renderer.document.add_paragraph(style=renderer._style("MD Normal"))
        bp.paragraph_format.left_indent = Mm(6)
        bp.paragraph_format.first_line_indent = Mm(-6)
        for segment in segments:
            text = clean_xml_text(segment.text)
            if segment.href and getattr(renderer.config.citations, "hyperlink_doi_and_url", True):
                add_hyperlink(
                    bp,
                    text,
                    segment.href,
                    font_name=renderer.fonts.for_text(rtl=renderer._rtl_for_text(text)),
                    rtl=renderer._rtl_for_text(text),
                )
            else:
                run = bp.add_run(text)
                renderer._configure_run(run, text)
        renderer._apply_text_policy(bp, " ".join(segment.text for segment in segments))
        name = f"cite_{ReferenceRegistry.safe_bookmark(key)}"[:40]
        if name not in renderer._bookmark_names:
            renderer._bookmark_names.add(name)
            add_bookmark(bp, name, renderer._bookmark_id)
            renderer._bookmark_id += 1
