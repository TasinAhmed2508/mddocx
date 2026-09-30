"""Text block surfaces: headings, paragraphs, quotes, callouts, and page flow.

Block surfaces are plain functions that receive the live ``DocxRenderer`` so they
share its document, numbering, reference registry, and text policy.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

from docx.enum.section import WD_SECTION
from docx.shared import Mm

from mddocx.ast.block import (
    BlockQuote,
    BulletList,
    Callout,
    DefinitionList,
    Heading,
    OrderedList,
    Paragraph,
)
from mddocx.ooxml.fields import add_bookmark, add_hidden_field, bookmark_name
from mddocx.ooxml.text import clean_xml_text
from mddocx.ooxml.utils import set_paragraph_bottom_border

from . import lists, sections

if TYPE_CHECKING:
    from mddocx.render.renderer import DocxRenderer


def render_heading(renderer: DocxRenderer, node: Heading) -> None:
    p = renderer.document.add_paragraph(style=renderer._style(f"MD Heading {node.level}"))
    renderer._apply_planned_paragraph_layout(p, node)
    if (
        renderer._heading_num_id is not None
        and node.level <= renderer.config.heading_numbering.max_level
    ):
        renderer.numbering.apply(p, renderer._heading_num_id, node.level - 1)
    renderer._render_inlines(p, node.children)
    renderer._apply_text_direction(p._p, renderer._plain_inline_text(node.children))
    if renderer.config.heading_bookmarks:
        plain = renderer._plain_inline_text(node.children) or f"Heading {renderer._bookmark_id}"
        target = (
            renderer.references.get(node.identifier)
            if node.identifier and renderer.references
            else None
        )
        name = target.bookmark if target else bookmark_name(plain, renderer._bookmark_id)
        base = name
        n = 2
        while name in renderer._bookmark_names:
            suffix = f"_{n}"
            name = base[: 40 - len(suffix)] + suffix
            n += 1
        renderer._bookmark_names.add(name)
        add_bookmark(p, name, renderer._bookmark_id)
        renderer._bookmark_id += 1
    if (
        renderer.config.references.enabled
        and node.level == 1
        and (
            renderer.config.references.equation_number_format == "section"
            or renderer.config.references.caption_number_format == "section"
        )
    ):
        counter_p = renderer.document.add_paragraph(style=renderer._style("MD Normal"))
        counter_p.paragraph_format.space_before = 0
        counter_p.paragraph_format.space_after = 0
        counter_p.paragraph_format.line_spacing = 0.01
        counter_p.paragraph_format.keep_with_next = True
        add_hidden_field(counter_p, r" SEQ Section \* ARABIC ", "1")
        if renderer.config.references.equation_number_format == "section":
            add_hidden_field(counter_p, r" SEQ Equation \r 0 ", "0")
        if renderer.config.references.caption_number_format == "section":
            for label in ("Figure", "Table", "Listing"):
                add_hidden_field(counter_p, f" SEQ {label} \r 0 ", "0")


def render_paragraph(renderer: DocxRenderer, node: Paragraph) -> None:
    p = renderer.document.add_paragraph(style=renderer._style("MD Normal"))
    renderer._render_inlines(p, node.children)
    renderer._apply_text_direction(p._p, renderer._plain_inline_text(node.children))


def render_block_quote(renderer: DocxRenderer, node: BlockQuote) -> None:
    for child in node.children:
        before = len(renderer.document.paragraphs)
        renderer._render_block(child)
        for p in renderer.document.paragraphs[before:]:
            p.style = renderer._style("MD Quote")


def render_callout(renderer: DocxRenderer, node: Callout) -> None:
    kind = (node.kind or "note").lower()
    label = renderer.config.callouts.labels.get(kind, kind.title())
    style_name = f"MD Callout {kind.title()}"
    first_paragraph = True
    for child in node.children or [Paragraph(children=[])]:
        if isinstance(child, Paragraph):
            p = renderer.document.add_paragraph(
                style=renderer._style(style_name)
                if style_name in renderer.document.styles
                else renderer._style("MD Quote")
            )
            if first_paragraph:
                r = p.add_run(
                    clean_xml_text(node.title or label) + (" — " if child.children else "")
                )
                r.bold = True
            renderer._render_inlines(p, child.children)
            renderer._apply_text_direction(p._p, renderer._plain_inline_text(child.children))
            first_paragraph = False
        elif isinstance(child, (BulletList, OrderedList)):
            if first_paragraph:
                p = renderer.document.add_paragraph(
                    style=renderer._style(style_name)
                    if style_name in renderer.document.styles
                    else renderer._style("MD Quote")
                )
                r = p.add_run(clean_xml_text(node.title or label))
                r.bold = True
                first_paragraph = False
            lists.render_list(renderer, child)
        else:
            renderer._render_block(child)


def render_definition_list(renderer: DocxRenderer, node: DefinitionList) -> None:
    for item in node.items:
        tp = renderer.document.add_paragraph(style=renderer._style("MD Normal"))
        tp.paragraph_format.keep_with_next = True
        renderer._render_inlines(tp, item.term, bold=True)
        dp = renderer.document.add_paragraph(style=renderer._style("MD Normal"))
        dp.paragraph_format.left_indent = Mm(8)
        renderer._render_inlines(dp, item.definition)


def render_horizontal_rule(renderer: DocxRenderer) -> None:
    p = renderer.document.add_paragraph(style=renderer._style("MD Normal"))
    set_paragraph_bottom_border(p)


def render_page_break(renderer: DocxRenderer) -> None:
    renderer.document.add_page_break()


def render_section_break(renderer: DocxRenderer) -> None:
    sec = renderer.document.add_section(WD_SECTION.NEW_PAGE)
    sections.configure_section(
        renderer,
        sec,
        preserve_page=bool(
            renderer.config.template and renderer.config.preserve_template_page_setup
        ),
    )
