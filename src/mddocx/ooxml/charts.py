from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from datetime import datetime
import zipfile

from lxml import etree
from openpyxl import Workbook

from mddocx.diagnostics import Diagnostic, MddocxError

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
CHART_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart"
PACKAGE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/package"

NS = {"w": W, "r": R, "pr": REL, "ct": CT}


@dataclass(slots=True)
class ChartEntry:
    token: str
    chart_type: str
    title: str
    categories: list[object]
    series: list[tuple[str, list[object]]]
    width_mm: float = 150.0
    height_mm: float = 85.0
    alt_text: str = "Chart"
    x_axis_title: str | None = None
    y_axis_title: str | None = None
    x_min: float | None = None
    x_max: float | None = None
    y_min: float | None = None
    y_max: float | None = None
    x_number_format: str | None = None
    y_number_format: str | None = None
    legend_position: str = "right"
    data_labels: bool = False
    show_gridlines: bool = True
    secondary_series: tuple[str, ...] = ()
    secondary_axis_title: str | None = None
    secondary_min: float | None = None
    secondary_max: float | None = None
    secondary_number_format: str | None = None
    style: int | None = None


def inject_charts(blob: bytes, entries: list[ChartEntry]) -> bytes:
    if not entries:
        return blob
    src = BytesIO(blob)
    with zipfile.ZipFile(src, "r") as zin:
        parts = {name: zin.read(name) for name in zin.namelist()}
    document = etree.fromstring(parts["word/document.xml"])
    rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
    content_types = etree.fromstring(parts["[Content_Types].xml"])
    next_rid = _next_rid(rels)
    next_docpr = _next_docpr(document)

    for index, entry in enumerate(entries, start=1):
        chart_name = f"chart{index}.xml"
        workbook_name = f"Microsoft_Excel_Worksheet{index}.xlsx"
        rid = f"rId{next_rid}"
        next_rid += 1
        rel = etree.SubElement(rels, f"{{{REL}}}Relationship")
        rel.set("Id", rid)
        rel.set("Type", CHART_REL)
        rel.set("Target", f"charts/{chart_name}")
        paragraph = _find_placeholder_paragraph(document, entry.token)
        if paragraph is None:
            raise MddocxError(
                Diagnostic("error", "CHART301", "Internal chart placeholder was not found.")
            )
        _replace_with_drawing(paragraph, rid, entry, next_docpr)
        next_docpr += 1
        parts[f"word/charts/{chart_name}"] = _chart_xml(entry, workbook_name)
        parts[f"word/charts/_rels/{chart_name}.rels"] = _chart_rels(workbook_name)
        parts[f"word/embeddings/{workbook_name}"] = _workbook_bytes(entry)
        _ensure_override(
            content_types,
            f"/word/charts/{chart_name}",
            "application/vnd.openxmlformats-officedocument.drawingml.chart+xml",
        )
    _ensure_default(
        content_types, "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    parts["word/document.xml"] = etree.tostring(
        document, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )
    parts["word/_rels/document.xml.rels"] = etree.tostring(
        rels, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )
    parts["[Content_Types].xml"] = etree.tostring(
        content_types, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )

    out = BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in parts.items():
            zout.writestr(name, data)
    return out.getvalue()


def _find_placeholder_paragraph(document, token: str):
    for p in document.xpath(".//w:p", namespaces={"w": W}):
        if token in "".join(p.itertext()):
            return p
    return None


def _replace_with_drawing(paragraph, rid: str, entry: ChartEntry, docpr_id: int) -> None:
    for child in list(paragraph):
        if child.tag != f"{{{W}}}pPr":
            paragraph.remove(child)
    ppr = paragraph.find(f"{{{W}}}pPr")
    if ppr is None:
        ppr = etree.Element(f"{{{W}}}pPr")
        paragraph.insert(0, ppr)
    jc = ppr.find(f"{{{W}}}jc")
    if jc is None:
        jc = etree.SubElement(ppr, f"{{{W}}}jc")
    jc.set(f"{{{W}}}val", "center")
    run = etree.SubElement(paragraph, f"{{{W}}}r")
    drawing = etree.SubElement(run, f"{{{W}}}drawing")
    inline = etree.SubElement(drawing, f"{{{WP}}}inline", nsmap={"wp": WP, "a": A, "c": C, "r": R})
    inline.set("distT", "0")
    inline.set("distB", "0")
    inline.set("distL", "0")
    inline.set("distR", "0")
    cx = str(int(entry.width_mm * 36000))
    cy = str(int(entry.height_mm * 36000))
    extent = etree.SubElement(inline, f"{{{WP}}}extent")
    extent.set("cx", cx)
    extent.set("cy", cy)
    effect = etree.SubElement(inline, f"{{{WP}}}effectExtent")
    for key in ("l", "t", "r", "b"):
        effect.set(key, "0")
    docpr = etree.SubElement(inline, f"{{{WP}}}docPr")
    docpr.set("id", str(docpr_id))
    docpr.set("name", f"Chart {docpr_id}")
    docpr.set("descr", entry.alt_text or entry.title or "Chart")
    cnv = etree.SubElement(inline, f"{{{WP}}}cNvGraphicFramePr")
    locks = etree.SubElement(cnv, f"{{{A}}}graphicFrameLocks")
    locks.set("noChangeAspect", "1")
    graphic = etree.SubElement(inline, f"{{{A}}}graphic")
    data = etree.SubElement(graphic, f"{{{A}}}graphicData")
    data.set("uri", C)
    chart = etree.SubElement(data, f"{{{C}}}chart")
    chart.set(f"{{{R}}}id", rid)


def _workbook_bytes(entry: ChartEntry) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(1, 1, "Category")
    for col, (name, _values) in enumerate(entry.series, start=2):
        ws.cell(1, col, name)
    rows = max([len(entry.categories), *(len(values) for _name, values in entry.series)] or [0])
    for row in range(rows):
        ws.cell(row + 2, 1, entry.categories[row] if row < len(entry.categories) else "")
        for col, (_name, values) in enumerate(entry.series, start=2):
            ws.cell(row + 2, col, values[row] if row < len(values) else None)
    fixed = datetime(2000, 1, 1, 0, 0, 0)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    stream = BytesIO()
    wb.save(stream)
    return _normalize_nested_zip(stream.getvalue())


def _normalize_nested_zip(blob: bytes) -> bytes:
    source = BytesIO(blob)
    with zipfile.ZipFile(source, "r") as zin:
        entries = []
        for name in sorted(zin.namelist()):
            data = zin.read(name)
            if name == "docProps/core.xml":
                try:
                    root = etree.fromstring(data)
                    for node in root.xpath(
                        "//*[local-name()='created' or local-name()='modified']"
                    ):
                        node.text = "2000-01-01T00:00:00Z"
                    data = etree.tostring(root, xml_declaration=False, encoding="UTF-8")
                except etree.XMLSyntaxError:
                    pass
            entries.append((name, data))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in entries:
            info = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            zout.writestr(info, data)
    return out.getvalue()


def _chart_xml(entry: ChartEntry, workbook_name: str) -> bytes:
    nsmap = {"c": C, "a": A, "r": R}
    root = etree.Element(f"{{{C}}}chartSpace", nsmap=nsmap)
    etree.SubElement(root, f"{{{C}}}date1904").set("val", "0")
    etree.SubElement(root, f"{{{C}}}lang").set("val", "en-US")
    etree.SubElement(root, f"{{{C}}}roundedCorners").set("val", "0")
    if entry.style is not None:
        etree.SubElement(root, f"{{{C}}}style").set("val", str(max(1, min(48, int(entry.style)))))
    chart = etree.SubElement(root, f"{{{C}}}chart")
    if entry.title:
        _title(chart, entry.title)
    plot = etree.SubElement(chart, f"{{{C}}}plotArea")
    etree.SubElement(plot, f"{{{C}}}layout")
    kind = entry.chart_type.lower()
    secondary_names = set(entry.secondary_series)
    secondary_indexes = [
        i for i, (name, _values) in enumerate(entry.series) if name in secondary_names
    ]
    primary_indexes = [i for i in range(len(entry.series)) if i not in secondary_indexes]
    if not primary_indexes and secondary_indexes:
        primary_indexes, secondary_indexes = secondary_indexes, []

    if kind in {"bar", "column"}:
        chart_node = _new_bar_chart(plot, kind, entry, primary_indexes)
        _axes(plot, chart_node, entry)
        if secondary_indexes:
            secondary = _new_bar_chart(plot, kind, entry, secondary_indexes)
            _secondary_axes(plot, secondary, entry)
    elif kind == "line":
        chart_node = _new_line_chart(plot, entry, primary_indexes)
        _axes(plot, chart_node, entry)
        if secondary_indexes:
            secondary = _new_line_chart(plot, entry, secondary_indexes)
            _secondary_axes(plot, secondary, entry)
    elif kind == "pie":
        chart_node = etree.SubElement(plot, f"{{{C}}}pieChart")
        etree.SubElement(chart_node, f"{{{C}}}varyColors").set("val", "1")
        _category_series(chart_node, entry, indexes=[0])
        if entry.data_labels:
            _data_labels(chart_node, kind, show_percent=True)
    elif kind == "scatter":
        chart_node = etree.SubElement(plot, f"{{{C}}}scatterChart")
        etree.SubElement(chart_node, f"{{{C}}}scatterStyle").set("val", "lineMarker")
        etree.SubElement(chart_node, f"{{{C}}}varyColors").set("val", "0")
        _scatter_series(chart_node, entry)
        if entry.data_labels:
            _data_labels(chart_node, kind)
        _value_axes(plot, chart_node, entry)
    else:
        raise MddocxError(
            Diagnostic("error", "CHART202", f"Unsupported chart type: {entry.chart_type}")
        )

    legend_pos = _legend_code(entry.legend_position)
    if legend_pos is not None:
        legend = etree.SubElement(chart, f"{{{C}}}legend")
        etree.SubElement(legend, f"{{{C}}}legendPos").set("val", legend_pos)
        etree.SubElement(legend, f"{{{C}}}layout")
        etree.SubElement(legend, f"{{{C}}}overlay").set("val", "0")
    etree.SubElement(chart, f"{{{C}}}plotVisOnly").set("val", "1")
    etree.SubElement(chart, f"{{{C}}}dispBlanksAs").set("val", "gap")
    external = etree.SubElement(root, f"{{{C}}}externalData")
    external.set(f"{{{R}}}id", "rId1")
    etree.SubElement(external, f"{{{C}}}autoUpdate").set("val", "0")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")


def _new_bar_chart(plot, kind: str, entry: ChartEntry, indexes: list[int]):
    node = etree.SubElement(plot, f"{{{C}}}barChart")
    etree.SubElement(node, f"{{{C}}}barDir").set("val", "bar" if kind == "bar" else "col")
    etree.SubElement(node, f"{{{C}}}grouping").set("val", "clustered")
    etree.SubElement(node, f"{{{C}}}varyColors").set("val", "0")
    _category_series(node, entry, indexes=indexes)
    if entry.data_labels:
        _data_labels(node, kind)
    return node


def _new_line_chart(plot, entry: ChartEntry, indexes: list[int]):
    node = etree.SubElement(plot, f"{{{C}}}lineChart")
    etree.SubElement(node, f"{{{C}}}grouping").set("val", "standard")
    etree.SubElement(node, f"{{{C}}}varyColors").set("val", "0")
    _category_series(node, entry, line=True, indexes=indexes)
    if entry.data_labels:
        _data_labels(node, "line")
    return node


def _legend_code(value: str) -> str | None:
    return {
        "right": "r",
        "r": "r",
        "left": "l",
        "l": "l",
        "top": "t",
        "t": "t",
        "bottom": "b",
        "b": "b",
        "top_right": "tr",
        "tr": "tr",
        "none": None,
        "off": None,
    }.get((value or "right").lower(), "r")


def _data_labels(parent, kind: str, *, show_percent: bool = False) -> None:
    labels = etree.SubElement(parent, f"{{{C}}}dLbls")
    etree.SubElement(labels, f"{{{C}}}showLegendKey").set("val", "0")
    etree.SubElement(labels, f"{{{C}}}showVal").set("val", "1")
    etree.SubElement(labels, f"{{{C}}}showCatName").set("val", "1" if kind == "pie" else "0")
    etree.SubElement(labels, f"{{{C}}}showSerName").set("val", "0")
    etree.SubElement(labels, f"{{{C}}}showPercent").set("val", "1" if show_percent else "0")
    etree.SubElement(labels, f"{{{C}}}showBubbleSize").set("val", "0")


def _title(chart, text: str) -> None:
    title = etree.SubElement(chart, f"{{{C}}}title")
    tx = etree.SubElement(title, f"{{{C}}}tx")
    rich = etree.SubElement(tx, f"{{{C}}}rich")
    etree.SubElement(rich, f"{{{A}}}bodyPr")
    etree.SubElement(rich, f"{{{A}}}lstStyle")
    p = etree.SubElement(rich, f"{{{A}}}p")
    r = etree.SubElement(p, f"{{{A}}}r")
    etree.SubElement(r, f"{{{A}}}rPr", lang="en-US")
    etree.SubElement(r, f"{{{A}}}t").text = text
    etree.SubElement(title, f"{{{C}}}layout")
    etree.SubElement(title, f"{{{C}}}overlay").set("val", "0")


def _axis_title(axis, text: str | None) -> None:
    if not text:
        return
    _title(axis, text)


def _category_series(
    parent, entry: ChartEntry, line: bool = False, indexes: list[int] | None = None
) -> None:
    rows = len(entry.categories)
    indexes = list(range(len(entry.series))) if indexes is None else indexes
    for idx in indexes:
        name, values = entry.series[idx]
        ser = etree.SubElement(parent, f"{{{C}}}ser")
        etree.SubElement(ser, f"{{{C}}}idx").set("val", str(idx))
        etree.SubElement(ser, f"{{{C}}}order").set("val", str(idx))
        _series_name(ser, idx + 2, name)
        if line:
            marker = etree.SubElement(ser, f"{{{C}}}marker")
            etree.SubElement(marker, f"{{{C}}}symbol").set("val", "circle")
        _str_ref(
            etree.SubElement(ser, f"{{{C}}}cat"), "Sheet1!$A$2:$A$%d" % (rows + 1), entry.categories
        )
        _num_ref(etree.SubElement(ser, f"{{{C}}}val"), _formula(idx + 2, rows), values)
        if line:
            etree.SubElement(ser, f"{{{C}}}smooth").set("val", "0")


def _scatter_series(parent, entry: ChartEntry) -> None:
    rows = len(entry.categories)
    xvalues = entry.categories
    for idx, (name, values) in enumerate(entry.series):
        ser = etree.SubElement(parent, f"{{{C}}}ser")
        etree.SubElement(ser, f"{{{C}}}idx").set("val", str(idx))
        etree.SubElement(ser, f"{{{C}}}order").set("val", str(idx))
        _series_name(ser, idx + 2, name)
        marker = etree.SubElement(ser, f"{{{C}}}marker")
        etree.SubElement(marker, f"{{{C}}}symbol").set("val", "circle")
        _num_ref(etree.SubElement(ser, f"{{{C}}}xVal"), "Sheet1!$A$2:$A$%d" % (rows + 1), xvalues)
        _num_ref(etree.SubElement(ser, f"{{{C}}}yVal"), _formula(idx + 2, rows), values)
        etree.SubElement(ser, f"{{{C}}}smooth").set("val", "0")


def _series_name(ser, column: int, name: str) -> None:
    tx = etree.SubElement(ser, f"{{{C}}}tx")
    ref = etree.SubElement(tx, f"{{{C}}}strRef")
    etree.SubElement(ref, f"{{{C}}}f").text = f"Sheet1!${_col(column)}$1"
    cache = etree.SubElement(ref, f"{{{C}}}strCache")
    etree.SubElement(cache, f"{{{C}}}ptCount").set("val", "1")
    pt = etree.SubElement(cache, f"{{{C}}}pt")
    pt.set("idx", "0")
    etree.SubElement(pt, f"{{{C}}}v").text = name


def _str_ref(parent, formula: str, values: list[object]) -> None:
    ref = etree.SubElement(parent, f"{{{C}}}strRef")
    etree.SubElement(ref, f"{{{C}}}f").text = formula
    cache = etree.SubElement(ref, f"{{{C}}}strCache")
    etree.SubElement(cache, f"{{{C}}}ptCount").set("val", str(len(values)))
    for idx, value in enumerate(values):
        pt = etree.SubElement(cache, f"{{{C}}}pt")
        pt.set("idx", str(idx))
        etree.SubElement(pt, f"{{{C}}}v").text = "" if value is None else str(value)


def _num_ref(parent, formula: str, values: list[object]) -> None:
    ref = etree.SubElement(parent, f"{{{C}}}numRef")
    etree.SubElement(ref, f"{{{C}}}f").text = formula
    cache = etree.SubElement(ref, f"{{{C}}}numCache")
    etree.SubElement(cache, f"{{{C}}}formatCode").text = "General"
    etree.SubElement(cache, f"{{{C}}}ptCount").set("val", str(len(values)))
    for idx, value in enumerate(values):
        pt = etree.SubElement(cache, f"{{{C}}}pt")
        pt.set("idx", str(idx))
        try:
            text = str(float(value))
        except (TypeError, ValueError):
            text = "0"
        etree.SubElement(pt, f"{{{C}}}v").text = text


def _axes(plot, chart_node, entry: ChartEntry) -> None:
    _category_axis_pair(
        plot,
        chart_node,
        entry,
        cat_id="123456",
        val_id="123457",
        cat_pos="b",
        val_pos="l",
        delete_cat=False,
        secondary=False,
    )


def _secondary_axes(plot, chart_node, entry: ChartEntry) -> None:
    _category_axis_pair(
        plot,
        chart_node,
        entry,
        cat_id="123458",
        val_id="123459",
        cat_pos="t",
        val_pos="r",
        delete_cat=True,
        secondary=True,
    )


def _category_axis_pair(
    plot,
    chart_node,
    entry: ChartEntry,
    *,
    cat_id: str,
    val_id: str,
    cat_pos: str,
    val_pos: str,
    delete_cat: bool,
    secondary: bool,
) -> None:
    etree.SubElement(chart_node, f"{{{C}}}axId").set("val", cat_id)
    etree.SubElement(chart_node, f"{{{C}}}axId").set("val", val_id)
    cat = etree.SubElement(plot, f"{{{C}}}catAx")
    etree.SubElement(cat, f"{{{C}}}axId").set("val", cat_id)
    _scaling(cat)
    etree.SubElement(cat, f"{{{C}}}delete").set("val", "1" if delete_cat else "0")
    etree.SubElement(cat, f"{{{C}}}axPos").set("val", cat_pos)
    etree.SubElement(cat, f"{{{C}}}tickLblPos").set("val", "none" if delete_cat else "nextTo")
    etree.SubElement(cat, f"{{{C}}}crossAx").set("val", val_id)
    etree.SubElement(cat, f"{{{C}}}crosses").set("val", "autoZero")
    etree.SubElement(cat, f"{{{C}}}auto").set("val", "1")
    etree.SubElement(cat, f"{{{C}}}lblAlgn").set("val", "ctr")
    etree.SubElement(cat, f"{{{C}}}lblOffset").set("val", "100")
    if not secondary:
        _axis_title(cat, entry.x_axis_title)

    val = etree.SubElement(plot, f"{{{C}}}valAx")
    etree.SubElement(val, f"{{{C}}}axId").set("val", val_id)
    _scaling(
        val,
        minimum=entry.secondary_min if secondary else entry.y_min,
        maximum=entry.secondary_max if secondary else entry.y_max,
    )
    etree.SubElement(val, f"{{{C}}}delete").set("val", "0")
    etree.SubElement(val, f"{{{C}}}axPos").set("val", val_pos)
    if entry.show_gridlines and not secondary:
        etree.SubElement(val, f"{{{C}}}majorGridlines")
    number_format = entry.secondary_number_format if secondary else entry.y_number_format
    etree.SubElement(
        val,
        f"{{{C}}}numFmt",
        formatCode=number_format or "General",
        sourceLinked="0" if number_format else "1",
    )
    etree.SubElement(val, f"{{{C}}}tickLblPos").set("val", "nextTo")
    etree.SubElement(val, f"{{{C}}}crossAx").set("val", cat_id)
    etree.SubElement(val, f"{{{C}}}crosses").set("val", "max" if secondary else "autoZero")
    etree.SubElement(val, f"{{{C}}}crossBetween").set("val", "between")
    _axis_title(val, entry.secondary_axis_title if secondary else entry.y_axis_title)


def _value_axes(plot, chart_node, entry: ChartEntry) -> None:
    x_id, y_id = "223456", "223457"
    etree.SubElement(chart_node, f"{{{C}}}axId").set("val", x_id)
    etree.SubElement(chart_node, f"{{{C}}}axId").set("val", y_id)
    x = etree.SubElement(plot, f"{{{C}}}valAx")
    etree.SubElement(x, f"{{{C}}}axId").set("val", x_id)
    _scaling(x, minimum=entry.x_min, maximum=entry.x_max)
    etree.SubElement(x, f"{{{C}}}delete").set("val", "0")
    etree.SubElement(x, f"{{{C}}}axPos").set("val", "b")
    if entry.show_gridlines:
        etree.SubElement(x, f"{{{C}}}majorGridlines")
    etree.SubElement(
        x,
        f"{{{C}}}numFmt",
        formatCode=entry.x_number_format or "General",
        sourceLinked="0" if entry.x_number_format else "1",
    )
    etree.SubElement(x, f"{{{C}}}tickLblPos").set("val", "nextTo")
    etree.SubElement(x, f"{{{C}}}crossAx").set("val", y_id)
    etree.SubElement(x, f"{{{C}}}crosses").set("val", "autoZero")
    etree.SubElement(x, f"{{{C}}}crossBetween").set("val", "midCat")
    _axis_title(x, entry.x_axis_title)

    y = etree.SubElement(plot, f"{{{C}}}valAx")
    etree.SubElement(y, f"{{{C}}}axId").set("val", y_id)
    _scaling(y, minimum=entry.y_min, maximum=entry.y_max)
    etree.SubElement(y, f"{{{C}}}delete").set("val", "0")
    etree.SubElement(y, f"{{{C}}}axPos").set("val", "l")
    if entry.show_gridlines:
        etree.SubElement(y, f"{{{C}}}majorGridlines")
    etree.SubElement(
        y,
        f"{{{C}}}numFmt",
        formatCode=entry.y_number_format or "General",
        sourceLinked="0" if entry.y_number_format else "1",
    )
    etree.SubElement(y, f"{{{C}}}tickLblPos").set("val", "nextTo")
    etree.SubElement(y, f"{{{C}}}crossAx").set("val", x_id)
    etree.SubElement(y, f"{{{C}}}crosses").set("val", "autoZero")
    etree.SubElement(y, f"{{{C}}}crossBetween").set("val", "midCat")
    _axis_title(y, entry.y_axis_title)


def _scaling(axis, *, minimum: float | None = None, maximum: float | None = None) -> None:
    scaling = etree.SubElement(axis, f"{{{C}}}scaling")
    etree.SubElement(scaling, f"{{{C}}}orientation").set("val", "minMax")
    if maximum is not None:
        etree.SubElement(scaling, f"{{{C}}}max").set("val", str(float(maximum)))
    if minimum is not None:
        etree.SubElement(scaling, f"{{{C}}}min").set("val", str(float(minimum)))


def _formula(column: int, rows: int) -> str:
    col = _col(column)
    return f"Sheet1!${col}$2:${col}${rows + 1}"


def _col(n: int) -> str:
    out = ""
    while n:
        n, rem = divmod(n - 1, 26)
        out = chr(65 + rem) + out
    return out


def _chart_rels(workbook_name: str) -> bytes:
    root = etree.Element("Relationships", nsmap={None: REL})
    rel = etree.SubElement(root, f"{{{REL}}}Relationship")
    rel.set("Id", "rId1")
    rel.set("Type", PACKAGE_REL)
    rel.set("Target", f"../embeddings/{workbook_name}")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")


def _next_rid(rels) -> int:
    values = []
    for rel in rels:
        rid = rel.get("Id", "")
        if rid.startswith("rId") and rid[3:].isdigit():
            values.append(int(rid[3:]))
    return max(values, default=0) + 1


def _next_docpr(document) -> int:
    ids = []
    for node in document.xpath(".//*[local-name()='docPr']"):
        try:
            ids.append(int(node.get("id", "0")))
        except ValueError:
            pass
    return max(ids, default=0) + 1


def _ensure_override(root, part_name: str, content_type: str) -> None:
    if root.xpath(f"./ct:Override[@PartName='{part_name}']", namespaces={"ct": CT}):
        return
    node = etree.SubElement(root, f"{{{CT}}}Override")
    node.set("PartName", part_name)
    node.set("ContentType", content_type)


def _ensure_default(root, extension: str, content_type: str) -> None:
    if root.xpath(f"./ct:Default[@Extension='{extension}']", namespaces={"ct": CT}):
        return
    node = etree.SubElement(root, f"{{{CT}}}Default")
    node.set("Extension", extension)
    node.set("ContentType", content_type)
