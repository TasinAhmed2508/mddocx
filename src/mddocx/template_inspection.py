from __future__ import annotations

from dataclasses import asdict, dataclass
import json
from pathlib import Path

from docx import Document


@dataclass(slots=True)
class TemplateInspection:
    path: str
    sections: int
    page_sizes: list[str]
    margins_mm: list[dict[str, float]]
    styles: int
    paragraph_styles: int
    character_styles: int
    table_styles: int
    custom_styles: list[str]
    header_text: list[str]
    footer_text: list[str]
    title: str
    author: str

    def to_json(self) -> str:
        return json.dumps(asdict(self), indent=2, ensure_ascii=False, sort_keys=True)

    def to_text(self) -> str:
        lines = [
            f"Template: {self.path}",
            f"Sections: {self.sections}",
            "Page sizes: " + ", ".join(self.page_sizes),
            f"Styles: {self.styles} (paragraph {self.paragraph_styles}, character {self.character_styles}, table {self.table_styles})",
            f"Custom styles: {len(self.custom_styles)}",
            f"Header text: {' | '.join(x for x in self.header_text if x) or '(none)'}",
            f"Footer text: {' | '.join(x for x in self.footer_text if x) or '(none)'}",
            f"Core title: {self.title or '(none)'}",
            f"Core author: {self.author or '(none)'}",
        ]
        if self.custom_styles:
            lines.append("Custom style names: " + ", ".join(self.custom_styles[:30]))
        return "\n".join(lines)


def inspect_template(path: str | Path) -> TemplateInspection:
    source = Path(path)
    doc = Document(str(source))
    page_sizes: list[str] = []
    margins: list[dict[str, float]] = []
    headers: list[str] = []
    footers: list[str] = []
    for sec in doc.sections:
        w = float(sec.page_width.mm)
        h = float(sec.page_height.mm)
        if abs(w - 210) < 2 and abs(h - 297) < 2:
            page_sizes.append("A4 portrait")
        elif abs(w - 297) < 2 and abs(h - 210) < 2:
            page_sizes.append("A4 landscape")
        elif abs(w - 215.9) < 2 and abs(h - 279.4) < 2:
            page_sizes.append("Letter portrait")
        elif abs(w - 279.4) < 2 and abs(h - 215.9) < 2:
            page_sizes.append("Letter landscape")
        else:
            page_sizes.append(f"{w:.1f}×{h:.1f} mm")
        margins.append(
            {
                "top": round(float(sec.top_margin.mm), 1),
                "bottom": round(float(sec.bottom_margin.mm), 1),
                "left": round(float(sec.left_margin.mm), 1),
                "right": round(float(sec.right_margin.mm), 1),
            }
        )
        headers.append(" ".join(p.text for p in sec.header.paragraphs).strip())
        footers.append(" ".join(p.text for p in sec.footer.paragraphs).strip())
    styles = list(doc.styles)
    paragraph = sum(1 for s in styles if int(s.type) == 1)
    character = sum(1 for s in styles if int(s.type) == 2)
    table = sum(1 for s in styles if int(s.type) == 3)
    custom = sorted(s.name for s in styles if not s.builtin)
    props = doc.core_properties
    return TemplateInspection(
        path=str(source),
        sections=len(doc.sections),
        page_sizes=page_sizes,
        margins_mm=margins,
        styles=len(styles),
        paragraph_styles=paragraph,
        character_styles=character,
        table_styles=table,
        custom_styles=custom,
        header_text=headers,
        footer_text=footers,
        title=props.title or "",
        author=props.author or "",
    )
