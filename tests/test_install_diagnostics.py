"""Install-environment diagnostics: package conflicts, math engines, Bangla fonts.

These cover the failure modes a new user meets right after ``pip install``:
another project providing the same ``mddocx`` package, a MathJax sidecar with no
JavaScript runtime, and a Bangla font that is configured but not installed.
"""

from __future__ import annotations

from importlib import metadata
from io import BytesIO

import pytest
from docx import Document

from mddocx.api import MarkdownWord
from mddocx.cli import _bangla_font_choice, main, warn_about_import_conflict
from mddocx.config import DEFAULT_BENGALI_FONT, FontConfig, RenderConfig
from mddocx.doctor import (
    DoctorReport,
    conflicting_distributions,
    install_conflict_message,
)
from mddocx.math.sidecar import MathJaxSidecar, MathJaxSidecarError, node_executable
from mddocx.ooxml import bangla
from mddocx.ooxml.bangla import (
    apply_bangla_font,
    is_font_installed,
    resolve_bangla_font,
)

_REQUIRED_DEPENDENCIES = (
    "markdown-it-py",
    "python-docx",
    "lxml",
    "Pillow",
    "Pygments",
    "PyYAML",
    "latex2mathml",
    "openpyxl",
    "prompt-toolkit",
)


@pytest.fixture(autouse=True)
def _clear_environment_caches():
    node_executable.cache_clear()
    bangla.installed_font_stems.cache_clear()
    yield
    node_executable.cache_clear()
    bangla.installed_font_stems.cache_clear()


def _healthy_report(**overrides) -> DoctorReport:
    fields: dict[str, object] = {
        "python": "3.11.9",
        "platform": "Test 1.0",
        "machine": "x86_64",
        "mddocx_version": "1.4.1",
        "dependencies": {name: "1.0" for name in _REQUIRED_DEPENDENCIES},
        "temp_directory_writable": True,
        "math_engine": "mathjax-sidecar",
    }
    fields.update(overrides)
    return DoctorReport(**fields)  # type: ignore[arg-type]


# --------------------------------------------------------------------------- #
# Conflicting installs
# --------------------------------------------------------------------------- #


def test_conflicting_distribution_is_detected(monkeypatch):
    monkeypatch.setattr(
        metadata,
        "packages_distributions",
        lambda: {"mddocx": ["mddocx", "mddocx-native"]},
    )
    assert conflicting_distributions() == ("mddocx",)


def test_this_distribution_is_never_its_own_conflict(monkeypatch):
    monkeypatch.setattr(
        metadata,
        "packages_distributions",
        lambda: {"mddocx": ["mddocx-native", "mddocx_native"]},
    )
    assert conflicting_distributions() == ()


def test_broken_metadata_is_not_fatal(monkeypatch):
    def explode():
        raise RuntimeError("corrupt metadata")

    monkeypatch.setattr(metadata, "packages_distributions", explode)
    assert conflicting_distributions() == ()


def test_conflict_message_names_the_repair_commands():
    message = install_conflict_message(("mddocx",))
    assert "pip uninstall -y mddocx mddocx-native" in message
    assert "pip install --force-reinstall mddocx-native" in message


def test_cli_warns_once_with_the_repair_commands(capsys, monkeypatch):
    monkeypatch.setattr("mddocx.cli.conflicting_distributions", lambda: ("mddocx",))
    assert warn_about_import_conflict() is True
    captured = capsys.readouterr()
    assert "conflicting mddocx installation" in captured.err
    assert "pip uninstall" in captured.err
    assert captured.out == ""


def test_cli_is_silent_when_the_install_is_clean(capsys, monkeypatch):
    monkeypatch.setattr("mddocx.cli.conflicting_distributions", lambda: ())
    assert warn_about_import_conflict() is False
    assert capsys.readouterr().err == ""


# --------------------------------------------------------------------------- #
# Doctor reporting
# --------------------------------------------------------------------------- #


def test_doctor_passes_for_a_healthy_install():
    assert _healthy_report().ok is True


def test_doctor_fails_on_a_conflicting_install():
    report = _healthy_report(install_conflict=("mddocx",))
    assert report.ok is False
    assert "Install conflict detected" in report.to_text()
    assert report.to_dict()["install_conflict"] == ("mddocx",)


def test_doctor_fails_when_no_math_engine_works():
    report = _healthy_report(math_engine=None, math_engine_error="UnsupportedLatexError: boom")
    assert report.ok is False
    assert "unavailable (UnsupportedLatexError: boom)" in report.to_text()


def test_doctor_reports_the_engine_that_actually_ran():
    text = _healthy_report(math_engine="latex2mathml").to_text()
    assert "Math engine in use: latex2mathml" in text
    assert _healthy_report(math_engine="latex2mathml").mathjax_ready is False
    assert _healthy_report(math_engine="mathjax-sidecar").mathjax_ready is True


def test_doctor_missing_engine_names_the_optional_extra():
    report = _healthy_report(
        dependencies={
            **_healthy_report().dependencies,
            "latex2mathml": "missing",
        }
    )
    assert report.ok is False


# --------------------------------------------------------------------------- #
# MathJax sidecar availability
# --------------------------------------------------------------------------- #


def test_script_sidecar_without_a_runtime_is_not_available(tmp_path, monkeypatch):
    script = tmp_path / "fake.mjs"
    script.write_text("// stand-in", encoding="utf-8")
    monkeypatch.delenv("MDDOCX_NODE", raising=False)
    monkeypatch.setattr("mddocx.math.sidecar.shutil.which", lambda name: None)
    node_executable.cache_clear()

    sidecar = MathJaxSidecar(script)
    assert sidecar.is_script is True
    assert sidecar.runtime is None
    assert sidecar.available is False
    assert "JavaScript runtime" in (sidecar.unavailable_reason() or "")
    with pytest.raises(MathJaxSidecarError, match="JavaScript runtime"):
        sidecar.convert("x^2", False)


def test_script_sidecar_is_available_when_a_runtime_exists(tmp_path, monkeypatch):
    script = tmp_path / "fake.mjs"
    script.write_text("// stand-in", encoding="utf-8")
    monkeypatch.delenv("MDDOCX_NODE", raising=False)
    monkeypatch.setattr("mddocx.math.sidecar.shutil.which", lambda name: "/usr/bin/node")
    node_executable.cache_clear()

    sidecar = MathJaxSidecar(script)
    assert sidecar.runtime is not None
    assert sidecar.available is True
    assert sidecar.unavailable_reason() is None


def test_native_sidecar_needs_no_javascript_runtime(tmp_path):
    executable = tmp_path / "mddocx-mathjax.exe"
    executable.write_bytes(b"stand-in")
    sidecar = MathJaxSidecar(executable)
    assert sidecar.is_script is False
    assert sidecar.available is True


def test_missing_sidecar_explains_itself(tmp_path):
    sidecar = MathJaxSidecar(tmp_path / "absent.mjs")
    assert sidecar.available is False
    assert "missing" in (sidecar.unavailable_reason() or "")


def test_configured_node_must_resolve(tmp_path, monkeypatch):
    monkeypatch.setenv("MDDOCX_NODE", str(tmp_path / "absent-node"))
    monkeypatch.setattr("mddocx.math.sidecar.shutil.which", lambda name: None)
    node_executable.cache_clear()
    assert node_executable() is None


# --------------------------------------------------------------------------- #
# Bangla font slot
# --------------------------------------------------------------------------- #


def test_bangla_font_choice_maps_flag_values():
    assert _bangla_font_choice(None) == DEFAULT_BENGALI_FONT
    assert _bangla_font_choice("   ") == DEFAULT_BENGALI_FONT
    assert _bangla_font_choice("none") is None
    assert _bangla_font_choice("Off") is None
    assert _bangla_font_choice("auto") == "auto"
    assert _bangla_font_choice("Nirmala UI") == "Nirmala UI"


def test_bangla_font_flag_reaches_the_render_config():
    from mddocx.cli import _config_from_args, build_parser

    def config(*extra: str):
        args = build_parser().parse_args(["report.md", *extra])
        return _config_from_args(args)

    assert config().fonts.bengali == DEFAULT_BENGALI_FONT
    assert config("--bengali-font", "auto").fonts.bengali == "auto"
    assert config("--bengali-font", "none").fonts.bengali is None
    assert config("--bengali-font", "Nirmala UI").fonts.bengali == "Nirmala UI"


def test_font_detection_matches_family_and_style_names(monkeypatch):
    monkeypatch.setattr(
        bangla, "installed_font_stems", lambda: frozenset({"sutonnymjregular", "nirmala", "vrinda"})
    )
    assert is_font_installed("SutonnyMJ") is True
    assert is_font_installed("Nirmala UI") is True
    assert is_font_installed("SolaimanLipi") is False
    assert is_font_installed("") is False


def test_missing_legacy_font_keeps_behaviour_and_explains_it(monkeypatch):
    monkeypatch.setattr(bangla, "installed_font_stems", lambda: frozenset())
    font, warning = resolve_bangla_font("SutonnyMJ")
    assert font == "SutonnyMJ"
    assert warning is not None
    assert "not installed" in warning
    assert "--bengali-font auto" in warning


def test_auto_prefers_the_legacy_font_when_installed(monkeypatch):
    monkeypatch.setattr(bangla, "installed_font_stems", lambda: frozenset({"sutonnymjregular"}))
    assert resolve_bangla_font("auto") == ("SutonnyMJ", None)


def test_auto_uses_an_installed_unicode_font(monkeypatch):
    monkeypatch.setattr(bangla, "installed_font_stems", lambda: frozenset({"nirmala"}))
    font, warning = resolve_bangla_font("auto")
    assert font == "Nirmala UI"
    assert warning is not None and "SutonnyMJ is not installed" in warning


def test_auto_without_any_bangla_font_disables_the_pass(monkeypatch):
    monkeypatch.setattr(bangla, "installed_font_stems", lambda: frozenset())
    font, warning = resolve_bangla_font("auto")
    assert font is None
    assert warning is not None and "no Unicode Bangla font" in warning


def test_document_without_bangla_text_skips_the_pass():
    blob = MarkdownWord().render_string("# English only\n\nPlain text.")
    assert apply_bangla_font(blob, DEFAULT_BENGALI_FONT) is blob


def test_auto_output_stays_unicode_when_no_bangla_font_is_installed(monkeypatch):
    from mddocx.compiler import Compiler

    monkeypatch.setattr(bangla, "installed_font_stems", lambda: frozenset())
    result = Compiler(RenderConfig(fonts=FontConfig(bengali="auto"))).compile_string(
        "**বাংলা English 123**"
    )
    assert result.success
    assert "FONT301" in {item.code for item in result.diagnostics}
    paragraph = Document(BytesIO(result.output_bytes)).paragraphs[0]
    assert paragraph.text == "বাংলা English 123"


def test_auto_output_uses_an_installed_unicode_font(monkeypatch):
    from mddocx.compiler import Compiler

    monkeypatch.setattr(bangla, "installed_font_stems", lambda: frozenset({"nirmala"}))
    result = Compiler(RenderConfig(fonts=FontConfig(bengali="auto"))).compile_string("বাংলা")
    assert result.success
    paragraph = Document(BytesIO(result.output_bytes)).paragraphs[0]
    assert paragraph.text == "বাংলা"
    assert paragraph.runs[0].font.name == "Nirmala UI"


def test_legacy_output_still_encodes_when_the_font_is_installed(monkeypatch):
    from mddocx.compiler import Compiler

    monkeypatch.setattr(bangla, "installed_font_stems", lambda: frozenset({"sutonnymjregular"}))
    result = Compiler(RenderConfig()).compile_string("**বাংলা English 123**")
    assert result.success
    assert {item.code for item in result.diagnostics} == set()
    paragraph = Document(BytesIO(result.output_bytes)).paragraphs[0]
    assert paragraph.text == "evsjv English 123"


# --------------------------------------------------------------------------- #
# First run without arguments
# --------------------------------------------------------------------------- #


def test_launcher_refuses_a_conflicting_install(capsys, monkeypatch):
    from mddocx_launcher import main as launcher_main

    monkeypatch.setattr("mddocx_launcher.conflicting_distributions", lambda: ("mddocx",))
    monkeypatch.delenv("MDDOCX_ALLOW_IMPORT_CONFLICT", raising=False)
    assert launcher_main(["--version"]) == 2
    captured = capsys.readouterr()
    assert captured.out == ""
    assert "not usable as installed" in captured.err
    assert "pip uninstall -y mddocx mddocx-native" in captured.err


def test_launcher_refuses_a_foreign_package(capsys, monkeypatch):
    from mddocx_launcher import main as launcher_main

    monkeypatch.setattr("mddocx_launcher.conflicting_distributions", lambda: ())
    monkeypatch.setattr(
        "mddocx_launcher.verify_installed_package",
        lambda: "the imported mddocx package is not this project (version 0.5.57)",
    )
    monkeypatch.delenv("MDDOCX_ALLOW_IMPORT_CONFLICT", raising=False)
    assert launcher_main(["--version"]) == 2
    assert "version 0.5.57" in capsys.readouterr().err


def test_launcher_continues_only_with_the_escape_hatch(capsys, monkeypatch):
    from mddocx_launcher import main as launcher_main

    monkeypatch.setattr("mddocx_launcher.conflicting_distributions", lambda: ("mddocx",))
    monkeypatch.setenv("MDDOCX_ALLOW_IMPORT_CONFLICT", "1")
    with pytest.raises(SystemExit) as exit_info:  # argparse exits for --version
        launcher_main(["--version"])
    assert exit_info.value.code == 0
    captured = capsys.readouterr()
    assert captured.out.strip().startswith("mddocx ")
    assert "MDDOCX_ALLOW_IMPORT_CONFLICT is set" in captured.err


def test_launcher_delegates_to_the_cli(capsys, monkeypatch):
    from mddocx_launcher import main as launcher_main

    monkeypatch.setattr("mddocx_launcher.conflicting_distributions", lambda: ())
    monkeypatch.setattr("mddocx_launcher.verify_installed_package", lambda: None)
    with pytest.raises(SystemExit) as exit_info:  # argparse exits for --version
        launcher_main(["--version"])
    assert exit_info.value.code == 0
    assert capsys.readouterr().out.strip().startswith("mddocx ")


def test_this_checkout_is_recognized_as_our_package():
    from mddocx_launcher import verify_installed_package

    assert verify_installed_package() is None


def test_foreign_package_is_detected(monkeypatch):
    import sys
    import types

    import mddocx_launcher

    foreign = types.ModuleType("mddocx")
    foreign.__version__ = "0.5.57"  # type: ignore[attr-defined]
    foreign.__author__ = "jw.zhou"  # type: ignore[attr-defined]
    monkeypatch.setitem(sys.modules, "mddocx", foreign)
    problem = mddocx_launcher.verify_installed_package()
    assert problem is not None
    assert "not this project" in problem
    assert "0.5.57" in problem


def test_warnings_reach_stderr_during_a_conversion(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr("mddocx.cli.conflicting_distributions", lambda: ())
    monkeypatch.setattr(bangla, "installed_font_stems", lambda: frozenset())
    source = tmp_path / "bangla.md"
    source.write_text("# শিরোনাম\n\nবাংলা লেখা\n", encoding="utf-8")

    assert main([str(source), "-o", str(tmp_path / "out.docx"), "--bengali-font", "auto"]) == 0
    captured = capsys.readouterr()
    assert "out.docx" in captured.out
    assert "FONT301" in captured.err
    assert "Remediation" in captured.err


def test_silent_conversion_writes_nothing_to_stderr(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr("mddocx.cli.conflicting_distributions", lambda: ())
    source = tmp_path / "plain.md"
    source.write_text("# Report\n\nLatin only.\n", encoding="utf-8")

    assert main([str(source), "-o", str(tmp_path / "out.docx")]) == 0
    captured = capsys.readouterr()
    assert captured.err == ""


def test_first_run_on_a_terminal_opens_the_guided_console(monkeypatch, capsys):
    import mddocx.cli as cli

    started: list[list[str]] = []
    monkeypatch.setattr(cli, "conflicting_distributions", lambda: ())
    monkeypatch.setattr(cli.sys.stdin, "isatty", lambda: True)
    monkeypatch.setattr(cli.sys.stdout, "isatty", lambda: True)
    monkeypatch.setattr(cli, "_main_shell", lambda argv: started.append(argv) or 0)

    assert cli.main([]) == 0
    assert started == [[]]
    out = capsys.readouterr().out
    assert "guided console" in out
    assert "mddocx report.md -o report.docx" in out


def test_doctor_reports_svg_support(monkeypatch):
    import mddocx.doctor as doctor

    monkeypatch.setattr(doctor, "svg_support_status", lambda: "unavailable: no cairo")
    report = doctor.run_doctor()
    assert report.svg_support == "unavailable: no cairo"
    assert "SVG conversion: unavailable: no cairo" in report.to_text()


def test_svg_support_reports_an_import_failure(monkeypatch):
    import builtins

    import mddocx.doctor as doctor

    real_import = builtins.__import__

    def failing_import(name, *args, **kwargs):
        if name == "cairosvg":
            raise OSError('no library called "cairo-2" was found')
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", failing_import)
    status = doctor.svg_support_status()
    assert status.startswith("unavailable")
    assert "cairo" in status


def test_first_run_without_a_terminal_keeps_the_usage_error(monkeypatch, capsys):
    monkeypatch.setattr("mddocx.cli.conflicting_distributions", lambda: ())
    with pytest.raises(SystemExit) as exit_info:
        main([])
    assert exit_info.value.code == 2
    assert "required: input" in capsys.readouterr().err
