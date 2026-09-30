"""Header and footer surfaces for a section.

Every header/footer paragraph is built from field templates and then handed to
``renderer._apply_text_policy`` so it inherits body fonts and RTL direction.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from mddocx.ooxml.fields import (
    add_num_pages,
    add_page_number,
    add_page_x_of_y,
    add_style_ref,
    render_field_template,
)

if TYPE_CHECKING:
    from mddocx.render.renderer import DocxRenderer


def configure_header_footer(renderer: DocxRenderer, sec) -> None:
    header_cfg = renderer.config.header
    footer_cfg = renderer.config.footer
    if header_cfg.different_first_page or footer_cfg.different_first_page:
        sec.different_first_page_header_footer = True
    if header_cfg.different_odd_even or footer_cfg.different_odd_even:
        settings = renderer.document.settings.element
        even = settings.find(qn("w:evenAndOddHeaders"))
        if even is None:
            even = OxmlElement("w:evenAndOddHeaders")
            settings.append(even)
        even.set(qn("w:val"), "true")

    def render_header(header, text: str | None, include_defaults: bool = True) -> None:
        header.is_linked_to_previous = False
        p = header.paragraphs[0]
        p.clear()
        pieces: list[str] = []
        if text:
            pieces.append(text)
        if include_defaults and header_cfg.document_title and renderer.config.title:
            pieces.append("{TITLE}")
        if pieces:
            render_field_template(
                p,
                " — ".join(pieces),
                {"TITLE": renderer.config.title or "", "AUTHOR": renderer.config.author or ""},
            )
        if include_defaults and header_cfg.section_title:
            if pieces:
                p.add_run(" — ")
            add_style_ref(p, renderer._style("MD Heading 1"))
        renderer._apply_text_policy(p, text)

    def render_footer(footer, text: str | None, include_defaults: bool = True) -> None:
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.clear()
        if text:
            render_field_template(
                p,
                text,
                {"TITLE": renderer.config.title or "", "AUTHOR": renderer.config.author or ""},
            )
        if not include_defaults:
            renderer._apply_text_policy(p, text)
            return
        has_text = bool(text)
        if footer_cfg.page_x_of_y:
            if has_text:
                p.add_run("  •  ")
            add_page_x_of_y(p)
        elif footer_cfg.page_number:
            if has_text:
                p.add_run("  •  ")
            add_page_number(p)
            if footer_cfg.num_pages:
                p.add_run(" / ")
                add_num_pages(p)
        elif footer_cfg.num_pages:
            if has_text:
                p.add_run("  •  ")
            add_num_pages(p)
        renderer._apply_text_policy(p, text)

    if (
        header_cfg.enabled
        or header_cfg.text
        or header_cfg.document_title
        or header_cfg.section_title
    ):
        render_header(sec.header, header_cfg.text)
    if header_cfg.different_first_page and header_cfg.first_page_text is not None:
        render_header(sec.first_page_header, header_cfg.first_page_text, False)
    if header_cfg.different_odd_even and header_cfg.even_page_text is not None:
        render_header(sec.even_page_header, header_cfg.even_page_text, False)

    if (
        footer_cfg.enabled
        or footer_cfg.text
        or footer_cfg.page_number
        or footer_cfg.num_pages
        or footer_cfg.page_x_of_y
    ):
        render_footer(sec.footer, footer_cfg.text)
    if footer_cfg.different_first_page and footer_cfg.first_page_text is not None:
        render_footer(sec.first_page_footer, footer_cfg.first_page_text, False)
    if footer_cfg.different_odd_even and footer_cfg.even_page_text is not None:
        render_footer(sec.even_page_footer, footer_cfg.even_page_text, False)
