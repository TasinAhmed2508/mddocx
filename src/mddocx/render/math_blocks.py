"""Math block surface: display equations and numbered equation tables."""

from __future__ import annotations

from typing import TYPE_CHECKING

from docx.enum.text import WD_ALIGN_PARAGRAPH

from mddocx.ast.block import MathBlock
from mddocx.ooxml.fields import add_chapter_seq_field, add_seq_field
from mddocx.ooxml.text import clean_xml_text
from mddocx.ooxml.utils import set_row_cant_split

from . import tables

if TYPE_CHECKING:
    from mddocx.render.renderer import DocxRenderer


def render_math_block(renderer: DocxRenderer, node: MathBlock) -> None:
    if (
        renderer.config.references.enabled
        and node.identifier
        and renderer.config.references.equation_numbering
    ):
        render_numbered_equation(renderer, node)
    else:
        p = renderer.document.add_paragraph(style=renderer._style("MD Equation"))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        renderer._apply_planned_paragraph_layout(p, node)
        renderer._append_math(p, node.source_text, display=True, source=node.source)
    if node.caption:
        cp = renderer.document.add_paragraph(style=renderer._style("MD Caption"))
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.add_run(clean_xml_text(node.caption))


def render_numbered_equation(renderer: DocxRenderer, node: MathBlock) -> None:
    table = renderer.document.add_table(rows=1, cols=2)
    table.autofit = False
    sec = renderer.document.sections[-1]
    usable = sec.page_width - sec.left_margin - sec.right_margin
    table.columns[0].width = int(usable * 0.88)
    table.columns[1].width = int(usable * 0.12)
    tables.remove_table_borders(table)
    p = table.cell(0, 0).paragraphs[0]
    p.style = renderer._style("MD Equation")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    renderer._append_math(p, node.source_text, display=True, source=node.source)
    np = table.cell(0, 1).paragraphs[0]
    np.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    target = renderer.references.get(node.identifier) if renderer.references else None
    np.add_run("(")
    if target and target.number is not None:
        if (
            renderer.config.references.equation_number_format == "section"
            and isinstance(target.number, str)
            and "." in target.number
        ):
            section_no, equation_no = target.number.split(".", 1)
            add_chapter_seq_field(
                np,
                "Equation",
                section_no,
                equation_no,
                target.bookmark,
                renderer._bookmark_id,
            )
        else:
            add_seq_field(
                np, "Equation", str(target.number), target.bookmark, renderer._bookmark_id
            )
        renderer._bookmark_id += 1
    else:
        add_seq_field(np, "Equation", "1")
    np.add_run(")")
    set_row_cant_split(table.rows[0], True)
