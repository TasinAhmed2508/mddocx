from __future__ import annotations

import json
from pathlib import Path
import zipfile

from docx import Document as WordDocument

from mddocx import MarkdownWord, MetadataConfig, RenderConfig
from mddocx.cli import main
from mddocx.project import build_project


def _document_text(path: Path) -> str:
    doc = WordDocument(path)
    return "\n".join(p.text for p in doc.paragraphs)


def test_render_removes_visible_ai_metadata_and_does_not_leak_export_identity_to_core_properties(tmp_path: Path):
    source = tmp_path / "chat.md"
    output = tmp_path / "chat.docx"
    source.write_text(
        """---
title: AI Conversation
author: AI Assistant
conversation_id: conv-99
model: model-x
created_at: 2026-08-16T10:00:00Z
---
Conversation ID: conv-99
Model: model-x
Exported at: 2026-08-16T10:05:00Z

# Actual answer

The useful content remains.
""",
        encoding="utf-8",
    )
    converter = MarkdownWord()
    converter.render_file(source, output)
    text = _document_text(output)
    assert "Conversation ID" not in text
    assert "Model: model-x" not in text
    assert "Exported at" not in text
    assert "Actual answer" in text
    doc = WordDocument(output)
    assert doc.core_properties.title == ""
    assert doc.core_properties.author == ""
    assert doc.core_properties.created is None
    assert doc.core_properties.modified is None
    assert any(d.code == "META101" for d in converter.diagnostics)


def test_keep_policy_retains_body_metadata_and_source_derived_core_properties(tmp_path: Path):
    source = tmp_path / "chat.md"
    output = tmp_path / "chat.docx"
    source.write_text(
        """---
title: AI Conversation
author: AI Assistant
conversation_id: conv-99
model: model-x
---
Conversation ID: conv-99
Model: model-x

# Answer
Text.
""",
        encoding="utf-8",
    )
    MarkdownWord(RenderConfig(metadata=MetadataConfig(ai_export="keep"))).render_file(source, output)
    text = _document_text(output)
    assert "Conversation ID: conv-99" in text
    assert "Model: model-x" in text
    doc = WordDocument(output)
    assert doc.core_properties.title == "AI Conversation"
    assert doc.core_properties.author == "AI Assistant"


def test_cli_metadata_inspect_and_clean(tmp_path: Path, capsys):
    source = tmp_path / "chat.md"
    cleaned = tmp_path / "clean.md"
    source.write_text(
        "Conversation ID: abc\nModel: test\nExported at: 2026-08-16T12:00:00Z\n\n# Answer\nHello.\n",
        encoding="utf-8",
    )
    assert main(["metadata", "inspect", str(source), "--json"]) == 0
    report = json.loads(capsys.readouterr().out)
    assert report["detected_export_metadata"] is True
    assert report["removed_lines"] >= 3
    assert main(["metadata", "clean", str(source), "-o", str(cleaned)]) == 0
    capsys.readouterr()
    text = cleaned.read_text(encoding="utf-8")
    assert "Conversation ID" not in text
    assert "# Answer" in text


def test_cli_ai_metadata_keep_flag(tmp_path: Path):
    source = tmp_path / "chat.md"
    output = tmp_path / "out.docx"
    source.write_text(
        "Conversation ID: abc\nModel: test\nExported at: 2026-08-16T12:00:00Z\n\n# Answer\nHello.\n",
        encoding="utf-8",
    )
    assert main([str(source), "-o", str(output), "--ai-metadata", "keep"]) == 0
    assert "Conversation ID: abc" in _document_text(output)


def test_project_build_sanitizes_each_source_before_merging(tmp_path: Path):
    root = tmp_path / "project"
    chapters = root / "chapters"
    chapters.mkdir(parents=True)
    (root / "mddocx.yml").write_text(
        """version: 1
output: build/document.docx
sources:
  - chapters/*.md
render:
  ai_metadata: auto
""",
        encoding="utf-8",
    )
    (chapters / "01-chat.md").write_text(
        """---
title: Export Title
conversation_id: c1
model: model-z
---
Conversation ID: c1
Model: model-z
Exported at: 2026-08-16T09:00:00Z

# Chapter One
Useful text.
""",
        encoding="utf-8",
    )
    result = build_project(root, force=True)
    text = _document_text(result.output_path)
    assert "Conversation ID" not in text
    assert "Model: model-z" not in text
    assert "Chapter One" in text
    doc = WordDocument(result.output_path)
    assert doc.core_properties.title != "Export Title"


def test_template_core_metadata_is_preserved_by_default(tmp_path: Path):
    template = tmp_path / "template.docx"
    source = tmp_path / "doc.md"
    output = tmp_path / "doc.docx"
    doc = WordDocument()
    doc.core_properties.author = "Corporate Template"
    doc.core_properties.title = "Template Title"
    doc.save(template)
    source.write_text("# Content\n\nHello.\n", encoding="utf-8")
    MarkdownWord(RenderConfig(template=template)).render_file(source, output)
    rendered = WordDocument(output)
    assert rendered.core_properties.author == "Corporate Template"
    assert rendered.core_properties.title == "Template Title"
